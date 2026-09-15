import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


template_path = Path(
    r"controlled_templates\8.1ADR REPORTING FORM.docx"
)

backup_path = Path(
    r"controlled_templates\8.1ADR REPORTING FORM.before_refinement.docx"
)

if not template_path.exists():
    raise FileNotFoundError(
        f"Template not found: {template_path.resolve()}"
    )

if not backup_path.exists():
    shutil.copy2(template_path, backup_path)

document = Document(template_path)


def set_cell_width(cell, width):
    cell.width = width

    cell_properties = cell._tc.get_or_add_tcPr()
    cell_width = cell_properties.find(qn("w:tcW"))

    if cell_width is None:
        cell_width = OxmlElement("w:tcW")
        cell_properties.append(cell_width)

    cell_width.set(qn("w:type"), "dxa")
    cell_width.set(qn("w:w"), str(width))


def format_table(table, proportions, font_size):
    section = document.sections[0]
    available_width = (
        section.page_width
        - section.left_margin
        - section.right_margin
    )

    total_proportion = sum(proportions)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for column_index, proportion in enumerate(proportions):
        column_width = int(
            available_width * proportion / total_proportion
        )

        for row in table.rows:
            if column_index < len(row.cells):
                set_cell_width(
                    row.cells[column_index],
                    column_width,
                )

    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(1)
                paragraph.paragraph_format.space_after = Pt(1)

                if row_index == 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

                for run in paragraph.runs:
                    run.font.size = Pt(font_size)

                    if row_index == 0:
                        run.bold = True


numbered_headings = (
    "1.0 ",
    "2.0 ",
    "3.0 ",
    "4.0 ",
    "5.0 ",
    "5.1 ",
    "5.2 ",
    "5.3 ",
    "5.4 ",
    "6.0 ",
    "6.1 ",
)

for paragraph in document.paragraphs:
    text = paragraph.text.strip()

    if text.startswith(numbered_headings):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Pt(0)
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(5)


if not document.tables:
    raise RuntimeError("Medical product table was not found.")

product_table = document.tables[0]

product_headers = [
    "Generic Name",
    "Brand Name",
    "Batch No",
    "Route, Dose, Frequency and Strength",
    "Date Started",
    "Date Stopped",
    "Indication",
    "Tick Suspected Medicine",
]

for cell, header in zip(product_table.rows[0].cells, product_headers):
    cell.text = header

format_table(
    product_table,
    proportions=[1.15, 1.0, 0.9, 1.65, 0.9, 0.9, 1.15, 0.55],
    font_size=8,
)


signature_table = None

for table in document.tables:
    if not table.rows:
        continue

    first_row_text = " ".join(
        cell.text.strip()
        for cell in table.rows[0].cells
    )

    if "Signatory" in first_row_text:
        signature_table = table
        break

if signature_table is not None:
    signature_table._element.getparent().remove(
        signature_table._element
    )

document.add_paragraph("")

heading = document.add_paragraph()
heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
heading.paragraph_format.space_before = Pt(8)
heading.paragraph_format.space_after = Pt(6)

heading_run = heading.add_run("Signatories")
heading_run.bold = True
heading_run.font.size = Pt(12)

signature_table = document.add_table(rows=1, cols=5)
signature_table.style = "Table Grid"

signature_headers = [
    "Signatory",
    "Name",
    "Designation",
    "Signature",
    "Date",
]

for cell, header in zip(
    signature_table.rows[0].cells,
    signature_headers,
):
    cell.text = header

signatories = [
    (
        "Prepared by",
        "AMEKO CHARLES",
        "DEPUTY Q.P.P.V.",
    ),
    (
        "Reviewed by",
        "YONA ODOI",
        "Q.P.P.V.",
    ),
    (
        "Authorised by",
        "KEITH ARUHO",
        "GROUP HEAD, REGULATORY AFFAIRS AND QUALITY",
    ),
]

for role, name, designation in signatories:
    row = signature_table.add_row()

    for cell, value in zip(
        row.cells,
        [role, name, designation, "", ""],
    ):
        cell.text = value

format_table(
    signature_table,
    proportions=[1.2, 1.35, 2.25, 1.05, 0.95],
    font_size=9,
)

for row in signature_table.rows[1:]:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(3)

document.save(template_path)

print("ADR template refinement completed successfully.")
print(f"Backup created at: {backup_path}")