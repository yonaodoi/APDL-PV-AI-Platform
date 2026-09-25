from datetime import date, datetime
from email.utils import parseaddr
from email.message import EmailMessage
from io import BytesIO
import secrets
from copy import deepcopy

from flask import (
    Blueprint,
    abort,
    flash,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from flask import current_app
from docx import Document

from app.db import query_all, query_one, transaction
from app.security import login_required
from app.services.case_completeness import refresh_case_completeness
from app.services.abacus_follow_up_pdf import build_abacus_follow_up_pdf
from app.services.abacus_follow_up_docx import build_abacus_follow_up_docx
from app.services.follow_up_email import (
    FollowUpEmailError,
    send_follow_up_email,
)
from app.services.gmail_oauth import (
    build_authorization_url,
    complete_authorization,
)


bp = Blueprint("case_review", __name__, url_prefix="/cases")


@bp.get("/email/gmail/connect")
@login_required
def connect_gmail():
    state = secrets.token_urlsafe(32)
    session["gmail_oauth_state"] = state
    try:
        authorization_url = build_authorization_url(
            current_app,
            state,
        )
    except RuntimeError as exc:
        flash(f"Gmail authorization failed: {exc}", "error")
        return redirect(url_for("case_review.follow_up_tasks"))
    return redirect(authorization_url)


@bp.get("/email/gmail/callback")
@login_required
def gmail_callback():
    state = session.pop("gmail_oauth_state", None)
    if not state or state != request.args.get("state"):
        flash("Gmail authorization could not be verified.", "error")
        return redirect(url_for("case_review.follow_up_tasks"))
    try:
        complete_authorization(
            current_app,
            request.url,
            state,
        )
    except Exception as exc:
        current_app.logger.exception("Gmail OAuth authorization failed")
        flash(f"Gmail authorization failed: {exc}", "error")
        return redirect(url_for("case_review.follow_up_tasks"))
    flash("Gmail is connected and ready to send follow-up forms.", "success")
    return redirect(url_for("case_review.follow_up_tasks"))


VALID_STATUSES = {
    "New",
    "Triage",
    "Follow-up requested",
    "Medical review",
    "Ready for submission",
    "Submitted",
    "Closed",
}


@bp.post("/<int:case_id>/review")
@login_required
def review_case(case_id):
    case = query_one(
        """
        SELECT case_id, case_number
        FROM pv.safety_cases
        WHERE case_id = %s
        """,
        (case_id,),
    )

    if case is None:
        abort(404)

    workflow_status = request.form.get("workflow_status", "")
    causality_assessment = request.form.get(
        "causality_assessment",
        "",
    ).strip()
    review_notes = request.form.get("review_notes", "").strip()
    follow_up_required = request.form.get("follow_up_required") == "on"
    follow_up_due_date = request.form.get("follow_up_due_date", "").strip()

    if workflow_status not in VALID_STATUSES:
        abort(400)

    try:
        parsed_follow_up_date = (
            datetime.strptime(follow_up_due_date, "%Y-%m-%d").date()
            if follow_up_due_date
            else None
        )
    except ValueError:
        flash("Follow-up due date is not valid.", "error")
        return redirect(url_for("cases.case_detail", case_id=case_id))

    with transaction() as cursor:
        cursor.execute(
            """
            UPDATE pv.safety_cases
            SET workflow_status = %s,
                causality_assessment = %s,
                follow_up_required = %s,
                follow_up_due_date = %s,
                updated_at = CURRENT_TIMESTAMP
            WHERE case_id = %s
            """,
            (
                workflow_status,
                causality_assessment or None,
                follow_up_required,
                parsed_follow_up_date,
                case_id,
            ),
        )

        details = (
            f"Status changed to {workflow_status}."
            + (f" Review note: {review_notes}" if review_notes else "")
        )

        cursor.execute(
            """
            INSERT INTO pv.case_audit_log (
                case_id,
                action,
                details,
                performed_by
            )
            VALUES (%s, %s, %s, %s)
            """,
            (
                case_id,
                "Case reviewed",
                details,
                session["user_id"],
            ),
        )

    updated_case = query_one(
        "SELECT * FROM pv.safety_cases WHERE case_id = %s",
        (case_id,),
    )
    products = query_all(
        "SELECT * FROM pv.case_products WHERE case_id = %s",
        (case_id,),
    )
    refresh_case_completeness(updated_case, products)

    flash(f"Case {case['case_number']} was updated.", "success")
    return redirect(url_for("cases.case_detail", case_id=case_id))


@bp.get("/follow-up-tasks")
@login_required
def follow_up_tasks():
    from app.services.case_follow_up import get_open_follow_up_tasks

    return render_template(
        "cases/follow_up_tasks.html",
        tasks=get_open_follow_up_tasks(),
        current_date=date.today(),
    )


@bp.get("/follow-up-reminders")
@login_required
def follow_up_reminders():
    from app.services.case_follow_up_reminders import get_open_reminders

    return render_template(
        "cases/follow_up_reminders.html",
        reminders=get_open_reminders(),
    )


@bp.post("/follow-up-reminders/<int:reminder_id>/acknowledge")
@login_required
def acknowledge_follow_up_reminder(reminder_id):
    from app.services.case_follow_up_reminders import acknowledge_reminder

    if acknowledge_reminder(reminder_id, session["user_id"]) is None:
        abort(404)
    flash("Reminder acknowledged.", "success")
    return redirect(url_for("case_review.follow_up_reminders"))


@bp.post("/follow-up-tasks/<int:task_id>/complete")
@login_required
def complete_follow_up_task(task_id):
    with transaction() as cursor:
        cursor.execute(
            """
            UPDATE pv.case_follow_up_tasks
            SET status = 'Completed',
                completed_at = CURRENT_TIMESTAMP,
                completed_by = %s,
                updated_at = CURRENT_TIMESTAMP
            WHERE task_id = %s
              AND status IN ('Open', 'In progress')
            RETURNING case_id
            """,
            (session["user_id"], task_id),
        )
        task = cursor.fetchone()

    if task is None:
        abort(404)

    flash("Follow-up task marked as completed.", "success")
    return redirect(url_for("case_review.follow_up_tasks"))


@bp.get("/follow-up-tasks/<int:task_id>/document")
@login_required
def download_follow_up_document(task_id):
    task = query_one(
        """
        SELECT tasks.*, cases.*
             , assigned_user.full_name AS assigned_to_name
        FROM pv.case_follow_up_tasks AS tasks
        JOIN pv.safety_cases AS cases ON cases.case_id = tasks.case_id
        WHERE tasks.task_id = %s
        """,
        (task_id,),
    )
    if task is None:
        abort(404)

    product = query_one(
        """
        SELECT *
        FROM pv.case_products
        WHERE case_id = %s
        ORDER BY case_product_id
        LIMIT 1
        """,
        (task["case_id"],),
    ) or {}
    checks = query_all(
        """
        SELECT check_code AS code, check_label AS label,
               status, message
        FROM pv.case_completeness_checks
        WHERE case_id = %s
          AND check_code = %s
          AND status = 'Review'
        ORDER BY check_id
        """,
        (task["case_id"], task["check_code"]),
    )

    pdf = build_abacus_follow_up_pdf(
        current_app.config["ABACUS_FOLLOW_UP_TEMPLATE_PATH"],
        task,
        product,
        task,
        checks,
    )
    filename = f"case-follow-up-{task['case_number']}-{task_id}.pdf"
    return send_file(
        pdf,
        as_attachment=True,
        download_name=filename,
        mimetype="application/pdf",
    )


@bp.get("/follow-up-tasks/<int:task_id>/document.docx")
@login_required
def download_follow_up_docx(task_id):
    task = query_one(
        """
        SELECT tasks.*, cases.*
        FROM pv.case_follow_up_tasks AS tasks
        JOIN pv.safety_cases AS cases ON cases.case_id = tasks.case_id
        LEFT JOIN pv.users AS assigned_user
            ON assigned_user.user_id = tasks.assigned_to
        WHERE tasks.task_id = %s
        """,
        (task_id,),
    )
    if task is None:
        abort(404)

    product = query_one(
        """
        SELECT *
        FROM pv.case_products
        WHERE case_id = %s
        ORDER BY case_product_id
        LIMIT 1
        """,
        (task["case_id"],),
    ) or {}
    check = query_one(
        """
        SELECT check_code AS code, check_label AS label,
               status, message
        FROM pv.case_completeness_checks
        WHERE case_id = %s
          AND check_code = %s
          AND status = 'Review'
        """,
        (task["case_id"], task["check_code"]),
    )
    if check is None:
        abort(404)

    document = build_abacus_follow_up_docx(
        task,
        product,
        task,
        check,
    )
    filename = f"case-follow-up-{task['case_number']}-{task_id}.docx"
    return send_file(
        document,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
    )


@bp.post("/follow-up-tasks/<int:task_id>/send-email")
@login_required
def send_follow_up_email_route(task_id):
    task = query_one(
        """
        SELECT tasks.*, cases.*, cases.reporter_email
        FROM pv.case_follow_up_tasks AS tasks
        JOIN pv.safety_cases AS cases ON cases.case_id = tasks.case_id
        WHERE tasks.task_id = %s
        """,
        (task_id,),
    )
    if task is None:
        abort(404)

    recipient = (task.get("reporter_email") or "").strip()
    if not recipient or parseaddr(recipient)[1] != recipient:
        flash(
            "This case does not have a valid reporter email address.",
            "error",
        )
        return redirect(url_for("case_review.follow_up_tasks"))


    product = query_one(
        """
        SELECT *
        FROM pv.case_products
        WHERE case_id = %s
        ORDER BY case_product_id
        LIMIT 1
        """,
        (task["case_id"],),
    ) or {}
    check = query_one(
        """
        SELECT check_code AS code, check_label AS label,
               status, message
        FROM pv.case_completeness_checks
        WHERE case_id = %s
          AND check_code = %s
          AND status = 'Review'
        """,
        (task["case_id"], task["check_code"]),
    )
    if check is None:
        abort(404)

    document = build_abacus_follow_up_docx(task, product, task, check)
    try:
        send_follow_up_email(current_app, recipient, task, document)
    except FollowUpEmailError as exc:
        with transaction() as cursor:
            cursor.execute(
                """
                INSERT INTO pv.case_follow_up_email_deliveries (
                    task_id, case_id, recipient_email, status,
                    error_message, sent_by
                )
                VALUES (%s, %s, %s, 'Failed', %s, %s)
                """,
                (
                    task_id,
                    task["case_id"],
                    recipient,
                    str(exc),
                    session["user_id"],
                ),
            )
        flash(str(exc), "error")
        return redirect(url_for("case_review.follow_up_tasks"))

    with transaction() as cursor:
        cursor.execute(
            """
            INSERT INTO pv.case_follow_up_email_deliveries (
                task_id, case_id, recipient_email, status, sent_by
            )
            VALUES (%s, %s, %s, 'Sent', %s)
            """,
            (task_id, task["case_id"], recipient, session["user_id"]),
        )
        cursor.execute(
            """
            INSERT INTO pv.case_audit_log (
                case_id, action, details, performed_by
            )
            VALUES (%s, %s, %s, %s)
            """,
            (
                task["case_id"],
                "Follow-up form emailed",
                f"Abacus Word follow-up form sent to {recipient}.",
                session["user_id"],
            ),
        )

    flash(f"Follow-up form sent to {recipient}.", "success")
    return redirect(url_for("case_review.follow_up_tasks"))


@bp.get("/follow-up-tasks/<int:task_id>/email-draft")
@login_required
def download_follow_up_email_draft(task_id):
    task = query_one(
        """
        SELECT tasks.*, cases.*
        FROM pv.case_follow_up_tasks AS tasks
        JOIN pv.safety_cases AS cases ON cases.case_id = tasks.case_id
        WHERE tasks.task_id = %s
        """,
        (task_id,),
    )
    if task is None:
        abort(404)

    recipient = (task.get("reporter_email") or "").strip()
    if not recipient or parseaddr(recipient)[1] != recipient:
        flash(
            "This case does not have a valid reporter email address.",
            "error",
        )
        return redirect(url_for("case_review.follow_up_tasks"))

    product = query_one(
        """
        SELECT *
        FROM pv.case_products
        WHERE case_id = %s
        ORDER BY case_product_id
        LIMIT 1
        """,
        (task["case_id"],),
    ) or {}
    check = query_one(
        """
        SELECT check_code AS code, check_label AS label,
               status, message
        FROM pv.case_completeness_checks
        WHERE case_id = %s
          AND check_code = %s
          AND status = 'Review'
        """,
        (task["case_id"], task["check_code"]),
    )
    if check is None:
        abort(404)

    document = build_abacus_follow_up_docx(task, product, task, check)
    message = EmailMessage()
    message["Subject"] = f"Case follow-up required - {task['case_number']}"
    message["From"] = current_app.config["SMTP_SENDER_EMAIL"]
    message["To"] = recipient
    message.set_content(
        "Please find attached the Abacus case follow-up form for "
        f"case {task['case_number']}. Please complete and return the form "
        "to the Pharmacovigilance team."
    )
    message.add_attachment(
        document.getvalue(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=f"case-follow-up-{task['case_number']}-{task['task_id']}.docx",
    )
    draft = BytesIO(message.as_bytes())
    draft.seek(0)
    return send_file(
        draft,
        as_attachment=True,
        download_name=f"case-follow-up-{task['case_number']}.eml",
        mimetype="message/rfc822",
    )


@bp.get("/follow-up-tasks/<int:task_id>/email.docx")
@login_required
def download_follow_up_email_docx(task_id):
    task = query_one(
        """
        SELECT tasks.*, cases.*
        FROM pv.case_follow_up_tasks AS tasks
        JOIN pv.safety_cases AS cases ON cases.case_id = tasks.case_id
        WHERE tasks.task_id = %s
        """,
        (task_id,),
    )
    if task is None:
        abort(404)

    recipient = (task.get("reporter_email") or "").strip()
    if not recipient or parseaddr(recipient)[1] != recipient:
        flash(
            "This case does not have a valid reporter email address.",
            "error",
        )
        return redirect(url_for("case_review.follow_up_tasks"))

    product = query_one(
        """
        SELECT *
        FROM pv.case_products
        WHERE case_id = %s
        ORDER BY case_product_id
        LIMIT 1
        """,
        (task["case_id"],),
    ) or {}
    check = query_one(
        """
        SELECT check_code AS code, check_label AS label,
               status, message
        FROM pv.case_completeness_checks
        WHERE case_id = %s
          AND check_code = %s
          AND status = 'Review'
        """,
        (task["case_id"], task["check_code"]),
    )
    if check is None:
        abort(404)

    form_document = build_abacus_follow_up_docx(
        task,
        product,
        task,
        check,
    )
    form_document.seek(0)
    form = Document(form_document)
    document = Document()
    document.add_heading("EMAIL MESSAGE", level=1)
    document.add_paragraph(f"To: {recipient}")
    document.add_paragraph(
        f"Subject: Case follow-up required - {task['case_number']}"
    )
    document.add_paragraph()
    document.add_paragraph(
        "Please find attached the Abacus case follow-up form for "
        f"case {task['case_number']}. Please complete and return the form "
        "to the Pharmacovigilance team."
    )
    document.add_page_break()
    for element in form.element.body:
        if element.tag.endswith("sectPr"):
            continue
        document.element.body.append(deepcopy(element))

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name=f"case-follow-up-email-{task['case_number']}.docx",
        mimetype=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
    )