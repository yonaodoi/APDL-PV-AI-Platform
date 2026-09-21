from pathlib import Path

from docx import Document
from pypdf import PdfReader


MAX_EXTRACTED_TEXT_LENGTH = 120000


def _clean_text(value):
    return " ".join((value or "").split())


def _extract_docx_text(file_path):
    document = Document(file_path)

    parts = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    for table in document.tables:
        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if values:
                parts.append(" | ".join(values))

    return "\n".join(parts)


def _extract_pdf_text(file_path):
    reader = PdfReader(file_path)
    parts = []

    for page in reader.pages:
        page_text = page.extract_text() or ""

        if page_text.strip():
            parts.append(page_text.strip())

    return "\n".join(parts)


def extract_reference_document_text(file_path):
    file_path = Path(file_path)
    extension = file_path.suffix.lower()

    if extension == ".pdf":
        text = _extract_pdf_text(file_path)

    elif extension == ".docx":
        text = _extract_docx_text(file_path)

    elif extension == ".doc":
        raise ValueError(
            "Legacy .doc files cannot be extracted automatically. "
            "Upload the document as PDF or DOCX."
        )

    else:
        raise ValueError(
            "Only PDF and DOCX reference documents can be extracted."
        )

    text = _clean_text(text)

    if not text:
        raise ValueError(
            "No readable text was found in the uploaded document."
        )

    return text[:MAX_EXTRACTED_TEXT_LENGTH]