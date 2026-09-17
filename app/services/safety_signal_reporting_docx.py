from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from app.services.ai_case_assessment_docx import (
    APDL_PURPLE,
    LIGHT_PURPLE,
    add_label_value_row,
    set_cell_shading,
    set_cell_text,
)


def _display_date(value):
    return value.strftime("%d %b %Y") if value else "Not recorded"


def _filter_text(filters):
    fields = (
        ("Product", filters.get("product")),
        ("Status", filters.get("status")),
        ("Priority", filters.get("priority")),
    )
    values = [
        f"{label}: {value}"
        for label, value in fields
        if value
    ]
    return " | ".join(values) if values else "No additional filters"


def build_safety_signal_reporting_docx(
    signals,
    report_filters,
    editable_content=None,
):
    editable_content = editable_content or {}
    control_values = editable_content.get("control_values") or []
    summary_values = editable_content.get("summary_values") or []
    section_titles = editable_content.get("section_titles") or []
    edited_rows = editable_content.get("signal_rows") or []
    edited_signatures = editable_content.get("signature_rows") or []

    def saved_value(values, index, fallback):
        if index < len(values) and values[index]:
            return values[index]
        return fallback

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(10)

    report_title = (
        editable_content.get("title")
        or "SAFETY SIGNAL REPORTING SUMMARY"
    )

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("ABACUS PARENTERAL DRUGS LIMITED")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string("666666")

    subtitle = header.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("REGULATORY AFFAIRS DEPARTMENT")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("666666")

    control_table = document.add_table(rows=4, cols=2)
    control_table.style = "Table Grid"

    add_label_value_row(
        control_table, 0, "TITLE",
        saved_value(control_values, 0, report_title),
    )
    add_label_value_row(
        control_table, 1, "REPORTING PERIOD",
        saved_value(
            control_values, 1,
            report_filters.get("reporting_period") or "All dates",
        ),
    )
    add_label_value_row(
        control_table, 2, "APPLIED FILTERS",
        saved_value(control_values, 2, _filter_text(report_filters)),
    )
    add_label_value_row(
        control_table, 3, "REPORT STATUS",
        saved_value(
            control_values, 3,
            "System-generated filtered report",
        ),
    )

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(16)
    title.paragraph_format.space_after = Pt(12)
    run = title.add_run(report_title)
    run.bold = True
    run.font.size = Pt(15)

    high_priority = sum(
        1 for signal in signals
        if signal["priority"] in ("High", "Critical")
    )
    under_evaluation = sum(
        1 for signal in signals
        if signal["status"] == "Under evaluation"
    )

    overview = document.add_table(rows=1, cols=3)
    overview.style = "Table Grid"
    values = (
        ("TOTAL SIGNALS", saved_value(summary_values, 0, len(signals))),
        ("HIGH / CRITICAL", saved_value(summary_values, 1, high_priority)),
        (
            "UNDER EVALUATION",
            saved_value(summary_values, 2, under_evaluation),
        ),
    )

    for index, (label, value) in enumerate(values):
        cell = overview.cell(0, index)
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)

        label_run = paragraph.add_run(f"{label}\n")
        label_run.bold = True
        label_run.font.size = Pt(8)
        label_run.font.color.rgb = RGBColor.from_string(APDL_PURPLE)

        value_run = paragraph.add_run(str(value))
        value_run.bold = True
        value_run.font.size = Pt(15)
        set_cell_shading(cell, LIGHT_PURPLE)

    document.add_paragraph()

    heading = document.add_paragraph()
    heading.paragraph_format.space_after = Pt(6)
    run = heading.add_run(
        saved_value(
            section_titles, 0,
            "1. FILTERED SAFETY SIGNAL REGISTER",
        )
    )
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(APDL_PURPLE)

    signal_table = document.add_table(rows=1, cols=8)
    signal_table.style = "Table Grid"

    headings = (
        "Signal ID", "Product", "Safety concern", "Source",
        "Owner", "Status", "Detected", "Priority",
    )

    for index, heading_text in enumerate(headings):
        set_cell_shading(signal_table.cell(0, index), APDL_PURPLE)
        set_cell_text(
            signal_table.cell(0, index), heading_text,
            bold=True, size=7, color="FFFFFF",
        )

    rows = edited_rows
    if not rows:
        rows = [
            [
                signal["signal_number"],
                signal["product_name"] or "Not recorded",
                signal["event_term"] or "Not recorded",
                signal["signal_source"] or "Not recorded",
                signal["owner_name"] or "Not recorded",
                signal["status"] or "Not recorded",
                _display_date(signal["date_detected"]),
                signal["priority"] or "Not recorded",
            ]
            for signal in signals
        ]

    if rows:
        for row in rows:
            cells = signal_table.add_row().cells
            for index in range(8):
                value = row[index] if index < len(row) else ""
                set_cell_text(cells[index], value, size=7)
    else:
        cells = signal_table.add_row().cells
        set_cell_text(
            cells[0],
            "No safety signals match the selected filters",
            size=8,
        )
        for index in range(1, 8):
            set_cell_text(cells[index], "", size=8)

    document.add_paragraph()

    signature_table = document.add_table(rows=4, cols=5)
    signature_table.style = "Table Grid"
    headings = (
        "Signatory", "Name", "Designation", "Signature", "Date",
    )

    for index, heading_text in enumerate(headings):
        set_cell_shading(signature_table.cell(0, index), APDL_PURPLE)
        set_cell_text(
            signature_table.cell(0, index), heading_text,
            bold=True, size=8, color="FFFFFF",
        )

    signatures = edited_signatures or (
        ("Prepared by", "AMEKO CHARLES", "DEPUTY Q.P.P.V.", "", ""),
        ("Reviewed by", "YONA ODOI", "Q.P.P.V.", "", ""),
        (
            "Authorised by", "KEITH ARUHO",
            "GROUP HEAD, RA & QUALITY", "", "",
        ),
    )

    for row_index, row in enumerate(signatures, start=1):
        values = list(row) + [""] * 5
        for column_index in range(5):
            set_cell_text(
                signature_table.cell(row_index, column_index),
                values[column_index], size=8,
            )

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(
        "System-generated safety signal reporting summary"
    )
    run.italic = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("666666")

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output