import re
from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


APDL_PURPLE = "5F3ED4"
LIGHT_PURPLE = "EEE9FF"
LIGHT_GREY = "E9E2E7"


def set_cell_shading(cell, fill):
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = cell_properties.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd"
    )

    if shading is None:
        from docx.oxml import OxmlElement

        shading = OxmlElement("w:shd")
        cell_properties.append(shading)

    shading.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill",
        fill,
    )


def set_cell_text(cell, text, bold=False, size=9, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)

    run = paragraph.add_run(str(text or ""))
    run.bold = bold
    run.font.size = Pt(size)

    if color:
        run.font.color.rgb = RGBColor.from_string(color)

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_label_value_row(table, row_index, label, value):
    label_cell = table.cell(row_index, 0)
    value_cell = table.cell(row_index, 1)

    set_cell_shading(label_cell, LIGHT_PURPLE)
    set_cell_text(label_cell, label, bold=True, size=8, color=APDL_PURPLE)
    set_cell_text(value_cell, value or "Not recorded", size=9)


def add_report_body(document, report_text):
    heading_pattern = re.compile(r"^\d+\.\s+.+")
    lines = (report_text or "").replace("*", "").splitlines()

    for line in lines:
        content = line.strip()

        if not content:
            continue

        if heading_pattern.match(content):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(5)

            run = paragraph.add_run(content)
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor.from_string(APDL_PURPLE)
            continue

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.15

        run = paragraph.add_run(content)
        run.font.size = Pt(10)


def build_ai_case_assessment_docx(case, product, report):
    document = Document()
    section = document.sections[0]

    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    for style_name in ("Normal",):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(10)

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_paragraph.paragraph_format.space_after = Pt(0)

    run = header_paragraph.add_run("ABACUS PARENTERAL DRUGS LIMITED")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string("666666")

    header_subtitle = header.add_paragraph()
    header_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_subtitle.paragraph_format.space_after = Pt(8)

    run = header_subtitle.add_run("REGULATORY AFFAIRS DEPARTMENT")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("666666")

    control_table = document.add_table(rows=4, cols=2)
    control_table.style = "Table Grid"
    control_table.autofit = False
    control_table.columns[0].width = Inches(1.6)
    control_table.columns[1].width = Inches(5.8)

    add_label_value_row(
        control_table,
        0,
        "TITLE",
        "AI CASE ASSESSMENT REPORT",
    )
    add_label_value_row(
        control_table,
        1,
        "DOCUMENT NO.",
        f"AI-CA/{case['case_number']}",
    )
    add_label_value_row(control_table, 2, "REVISION STATUS", "00")
    add_label_value_row(
        control_table,
        3,
        "REPORT STATUS",
        report.get("generation_status", "Generated"),
    )

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(16)
    title.paragraph_format.space_after = Pt(12)

    run = title.add_run("AI CASE ASSESSMENT REPORT")
    run.bold = True
    run.font.size = Pt(15)

    case_table = document.add_table(rows=4, cols=2)
    case_table.style = "Table Grid"
    case_table.autofit = False
    case_table.columns[0].width = Inches(1.8)
    case_table.columns[1].width = Inches(5.6)

    add_label_value_row(case_table, 0, "CASE NUMBER", case["case_number"])
    add_label_value_row(
        case_table,
        1,
        "PRODUCT",
        (product or {}).get("product_name", "Not recorded"),
    )
    add_label_value_row(
        case_table,
        2,
        "EVENT",
        case.get("event_description", "Not recorded"),
    )
    add_label_value_row(
        case_table,
        3,
        "GENERATED",
        report["created_at"].strftime("%d %B %Y %H:%M"),
    )

    document.add_paragraph()

    add_report_body(document, report["report_text"])

    document.add_paragraph()

    signature_table = document.add_table(rows=4, cols=5)
    signature_table.style = "Table Grid"
    signature_table.autofit = False

    column_widths = (
        Inches(1.05),
        Inches(1.2),
        Inches(2.25),
        Inches(1.2),
        Inches(1.15),
    )

    for index, width in enumerate(column_widths):
        for cell in signature_table.columns[index].cells:
            cell.width = width

    headings = (
        "Signatory",
        "Name",
        "Designation",
        "Signature",
        "Date",
    )

    for index, heading in enumerate(headings):
        set_cell_shading(signature_table.cell(0, index), APDL_PURPLE)
        set_cell_text(
            signature_table.cell(0, index),
            heading,
            bold=True,
            size=8,
            color="FFFFFF",
        )

    signature_rows = (
        (
            "Prepared by",
            report.get("prepared_by_name") or "AMEKO CHARLES",
            report.get("prepared_by_designation")
            or "DEPUTY Q.P.P.V.",
        ),
        (
            "Reviewed by",
            report.get("reviewed_by_name") or "YONA ODOI",
            report.get("reviewed_by_designation") or "Q.P.P.V.",
        ),
        (
            "Authorised by",
            report.get("authorised_by_name") or "KEITH ARUHO",
            report.get("authorised_by_designation")
            or "GROUP HEAD, RA & QUALITY",
        ),
    )

    for row_index, (role, name, designation) in enumerate(
        signature_rows,
        start=1,
    ):
        set_cell_text(signature_table.cell(row_index, 0), role, size=8)
        set_cell_text(signature_table.cell(row_index, 1), name, size=8)
        set_cell_text(
            signature_table.cell(row_index, 2),
            designation,
            size=8,
        )
        set_cell_text(signature_table.cell(row_index, 3), "", size=8)
        set_cell_text(signature_table.cell(row_index, 4), "", size=8)
    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = footer_paragraph.add_run(
        "AI-generated draft — QPPV/medical reviewer approval required"
    )
    run.italic = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("666666")

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output