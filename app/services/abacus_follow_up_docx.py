from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from app.services.ai_case_assessment_docx import (
    APDL_PURPLE,
    LIGHT_PURPLE,
    set_cell_shading,
    set_cell_text,
)


def _date(value):
    return value.strftime("%d/%m/%Y") if value else ""


def _value(value):
    return str(value) if value not in (None, "") else ""


FOLLOW_UP_QUESTIONS = {
    "minimum_patient": (
        "Please provide a unique patient identifier, such as the "
        "patient's initials, date of birth, age, or other information "
        "that allows this patient to be distinguished from other patients."
    ),
    "identifiable_reporter": (
        "Please provide the reporter's name and at least one contact "
        "detail, such as an email address, telephone number, or "
        "organisation."
    ),
    "suspected_product": (
        "Please confirm the suspected medicinal product name and provide "
        "any available product details, including strength, dosage form, "
        "batch or lot number, dose, route, and indication."
    ),
    "reported_event": (
        "Please provide a clear description of the adverse event or "
        "reaction, including the signs, symptoms, diagnosis, and relevant "
        "clinical details."
    ),
    "event_onset_date": (
        "Please provide the date on which the adverse event or reaction "
        "started. If the exact date is unknown, provide the best estimate "
        "and explain the basis for that estimate."
    ),
    "event_outcome": (
        "Please confirm the current outcome of the adverse event or "
        "reaction: recovered/resolved, recovering/resolving, not "
        "recovered/not resolved, recovered with sequelae, fatal, or "
        "unknown."
    ),
    "seriousness_basis": (
        "Please confirm the seriousness criterion that applies, such as "
        "death, life-threatening condition, hospitalisation, disability, "
        "congenital anomaly, or another medically important condition."
    ),
    "follow_up_due_date": (
        "Please confirm the date by which the requested follow-up "
        "information is expected."
    ),
}


def _add_heading(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(APDL_PURPLE)


def _add_label_value_table(document, rows):
    table = document.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    table.autofit = False
    for index, (label, value) in enumerate(rows):
        set_cell_shading(table.cell(index, 0), LIGHT_PURPLE)
        set_cell_text(
            table.cell(index, 0),
            label,
            bold=True,
            size=8,
            color=APDL_PURPLE,
        )
        set_cell_text(table.cell(index, 1), _value(value), size=9)
        table.cell(index, 0).width = Inches(2.1)
        table.cell(index, 1).width = Inches(4.7)
    return table


def build_abacus_follow_up_docx(case, product, task, check):
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(10)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("ABACUS PARENTERAL DRUGS LTD")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string("666666")
    subtitle = section.header.add_paragraph("REGULATORY AFFAIRS DEPARTMENT")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(9)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("CASE FOLLOW-UP FORM")
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor.from_string(APDL_PURPLE)

    control = _add_label_value_table(
        document,
        (
            ("DOCUMENT No", "SF/RA/012.2"),
            ("REVISION STATUS", "00"),
            ("EFFECTIVE DATE", "31/08/2026"),
            ("FOLLOW-UP DUE DATE", _date(task["due_date"])),
            ("ASSIGNED TO", task.get("assigned_to_name") or ""),
        ),
    )
    document.add_paragraph()

    _add_heading(document, "PART A: CASE IDENTIFICATION")
    patient_identifier = case.get("patient_initials") or ""
    patient_detail = patient_identifier
    if case.get("patient_date_of_birth"):
        patient_detail += f" | DOB: {_date(case['patient_date_of_birth'])}"
    elif case.get("patient_age_years") is not None:
        patient_detail += f" | Age: {case['patient_age_years']} years"

    _add_label_value_table(
        document,
        (
            ("APDL CASE NUMBER", case["case_number"]),
            ("INTERNAL CASE ID", case["case_id"]),
            ("PATIENT IDENTIFIER", patient_detail),
            ("INITIAL DATE RECEIVED (DAY 0)", _date(case["received_date"])),
            ("SUSPECTED MEDICINAL PRODUCT", product.get("product_name")),
            ("ADVERSE EVENT / REACTION", case.get("event_description")),
        ),
    )

    _add_heading(document, "PART B: REASON FOR FOLLOW-UP")
    reason_table = document.add_table(rows=1, cols=2)
    reason_table.style = "Table Grid"
    set_cell_shading(reason_table.cell(0, 0), LIGHT_PURPLE)
    set_cell_text(
        reason_table.cell(0, 0),
        f"[X] {check['label']}",
        bold=True,
        size=9,
        color=APDL_PURPLE,
    )
    set_cell_text(reason_table.cell(0, 1), check["message"], size=9)

    _add_heading(document, "PART C: FOLLOW-UP QUESTION")
    question = document.add_paragraph()
    question.paragraph_format.space_after = Pt(8)
    question_run = question.add_run(
        FOLLOW_UP_QUESTIONS.get(
            check["code"],
            (
                "Please provide the information requested under Part B "
                "and return this form to the APDL Regulatory Affairs "
                "Department."
            ),
        )
    )
    question_run.font.size = Pt(10)
    question_run.bold = True
    response_label = document.add_paragraph()
    response_label.paragraph_format.space_before = Pt(10)
    response_label.paragraph_format.space_after = Pt(4)
    response_label.add_run("Response:").bold = True
    document.add_paragraph("\n\n\n")

    document.add_page_break()
    _add_heading(document, "SIGNATORIES")
    signatory_table = document.add_table(rows=4, cols=5)
    signatory_table.style = "Table Grid"
    signatory_table.autofit = False
    column_widths = (
        Inches(1.05),
        Inches(1.2),
        Inches(2.25),
        Inches(1.2),
        Inches(1.15),
    )
    for index, width in enumerate(column_widths):
        for cell in signatory_table.columns[index].cells:
            cell.width = width

    signatory_headers = (
        "Signatory",
        "Name",
        "Designation",
        "Signature",
        "Date",
    )
    for column, label in enumerate(signatory_headers):
        set_cell_shading(signatory_table.cell(0, column), LIGHT_PURPLE)
        set_cell_text(
            signatory_table.cell(0, column),
            label,
            bold=True,
            size=8,
            color=APDL_PURPLE,
        )

    signatories = (
        ("Prepared by", "AMEKO CHARLES", "DEPUTY Q.P.P.V."),
        ("Reviewed by", "YONA ODOI", "Q.P.P.V."),
        (
            "Authorised by",
            "KEITH ARUHO",
            "GROUP HEAD, RA & QUALITY",
        ),
    )
    for row, (role, name, designation) in enumerate(signatories, 1):
        set_cell_text(signatory_table.cell(row, 0), role, size=8)
        set_cell_text(signatory_table.cell(row, 1), name, size=8)
        set_cell_text(signatory_table.cell(row, 2), designation, size=8)
        set_cell_text(signatory_table.cell(row, 3), "\n\n", size=8)
        set_cell_text(signatory_table.cell(row, 4), "\n", size=8)

    closing = document.add_paragraph()
    closing.paragraph_format.space_before = Pt(12)
    closing.add_run(
        "Signatures confirm review and authorisation of this follow-up "
        "request in accordance with the Abacus controlled form."
    ).font.size = Pt(8)

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output
