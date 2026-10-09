"""One follow-up request per safety case, built from its missing information.

The completeness checklist decides what is missing. Items the reporter can
answer become numbered questions on a single Word form; items that are the
PV team's own work (the follow-up plan, MedDRA coding) never leave the
building.
"""

from datetime import date
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from app.services.ai_case_assessment_docx import (
    APDL_PURPLE,
    LIGHT_PURPLE,
    set_cell_shading,
    set_cell_text,
)
from app.services import company_profile as _company_profile
from app.services.company_profile import DEFAULTS


def company():
    """The company profile (looked up each time, so changes apply at once)."""
    return _company_profile.company()


# Checks the PV team resolves itself; never sent to a reporter.
INTERNAL_CODES = {"follow_up_due_date", "event_coding"}
# Contradictions between recorded values; often our own typing errors, so the
# PV team checks them first. They can still be added to a request by hand.
CHECK_FIRST_CODES = {"fatal_outcome_consistency", "date_sequence", "patient_consistency"}

OUTCOME_OPTIONS = (
    "Recovered/resolved",
    "Recovering/resolving",
    "Not recovered/not resolved",
    "Recovered with sequelae",
    "Fatal",
    "Unknown",
)
SERIOUSNESS_OPTIONS = (
    "Resulted in death",
    "Life-threatening",
    "Required or prolonged hospitalisation",
    "Persistent or significant disability/incapacity",
    "Congenital anomaly/birth defect",
    "Other medically important condition",
)

# Where on the case edit form each item is filled in or corrected.
EDIT_SECTION_FOR_CODE = {
    "minimum_patient": "patient-information",
    "patient_consistency": "patient-information",
    "identifiable_reporter": "report-details",
    "suspected_product": "product-details",
    "reported_event": "event-assessment",
    "event_onset_date": "event-assessment",
    "event_outcome": "event-assessment",
    "seriousness_basis": "event-assessment",
    "fatal_outcome_consistency": "event-assessment",
    "date_sequence": "event-assessment",
}

# The exact form field to open for each item (its id on the case form).
EDIT_FIELD_FOR_CODE = {
    "minimum_patient": "patient_initials",
    "patient_consistency": "patient_date_of_birth",
    "identifiable_reporter": "reporter_name",
    "suspected_product": "product_name",
    "reported_event": "event_description",
    "event_onset_date": "event_onset_date",
    "event_outcome": "event_outcome",
    "seriousness_basis": "seriousness_criteria",
    "fatal_outcome_consistency": "event_outcome",
    "date_sequence": "event_onset_date",
}

# Plain-language headings for the reporter (the checklist labels are ours).
REPORTER_LABELS = {
    "minimum_patient": "Patient details",
    "identifiable_reporter": "Your contact details",
    "suspected_product": "Suspected medicine",
    "reported_event": "Description of the reaction",
    "event_onset_date": "When the reaction started",
    "event_outcome": "Outcome of the reaction",
    "seriousness_basis": "Why the reaction is serious",
    "fatal_outcome_consistency": "Outcome and seriousness",
    "date_sequence": "Dates",
    "patient_consistency": "Patient details",
}

def form_control():
    """Document-control block of the follow-up form (from the company profile)."""
    profile = company()
    return (
        ("DOCUMENT No", profile["form_code"]),
        ("REVISION STATUS", profile["form_revision"]),
        ("EFFECTIVE DATE", profile["form_effective_date"]),
    )


def form_signatories():
    from app.services.company_profile import signatory

    return signatory("reviewed"), signatory("authorised")


DEFAULT_PV_PHONE = DEFAULTS["pv_phone"]


def format_sent(value):
    """'08 Oct 2026 at 22:51 (UTC+03:00)' in this computer's time zone."""
    if not value:
        return ""
    if getattr(value, "tzinfo", None) is None:
        value = value.astimezone()
    else:
        value = value.astimezone()
    offset = value.strftime("%z")
    offset = f"UTC{offset[:3]}:{offset[3:]}" if offset else ""
    return value.strftime("%d %b %Y at %H:%M") + (f" ({offset})" if offset else "")


def _fmt(value):
    return value.strftime("%d %b %Y") if value else ""


def is_reporter_item(check):
    return check.get("status") == "Review" and check.get("code") not in INTERNAL_CODES


def _event(case):
    text = " ".join((case.get("event_description") or "").split())
    if not text:
        return "the reported reaction"
    return text if len(text) <= 80 else text[:79] + "…"


def build_request_items(case, product, checks):
    """Numbered questions for the reporter, from the case's open checks.

    Each item: code, label, question, answer ("text", "date", "options",
    "multi") and options where relevant.
    """
    product_name = (product or {}).get("product_name") or "the suspected medicine"
    event = _event(case)
    items = []
    for check in checks or []:
        if not is_reporter_item(check):
            continue
        code = check["code"]
        item = {
            "code": code,
            "label": REPORTER_LABELS.get(code) or check.get("label") or code,
            "answer": "text",
            "options": (),
        }

        if code == "minimum_patient":
            item["question"] = (
                "Please give at least one detail that identifies the patient: "
                "initials, age or date of birth, and sex. Do not send the "
                "patient's full name."
            )
        elif code == "identifiable_reporter":
            item["question"] = (
                "Please confirm your name, profession and a telephone number "
                "or email address where we can contact you about this report."
            )
        elif code == "suspected_product":
            item["question"] = (
                "Which medicine do you suspect caused the reaction? Please give "
                "the name, strength, batch number, dose, route, the reason it "
                "was given, and the start and stop dates."
            )
        elif code == "reported_event":
            item["question"] = (
                f"Please describe the reaction seen with {product_name}: the "
                "signs and symptoms, any diagnosis, and how it was treated."
            )
        elif code == "event_onset_date":
            item["question"] = (
                f"On what date did the reaction ({event}) start? If the exact "
                "date is not known, give your best estimate (for example the "
                "month and year)."
            )
            item["answer"] = "date"
        elif code == "event_outcome":
            item["question"] = (
                f"What is the patient's current condition regarding the "
                f"reaction ({event})? Please tick one."
            )
            item["answer"] = "options"
            item["options"] = OUTCOME_OPTIONS
        elif code == "seriousness_basis":
            item["question"] = (
                "The reaction has been reported as serious. Please tick every "
                "reason that applies and give brief details (for example "
                "hospital admission and discharge dates)."
            )
            item["answer"] = "multi"
            item["options"] = SERIOUSNESS_OPTIONS
        elif code in ("fatal_outcome_consistency", "date_sequence", "patient_consistency"):
            message = (check.get("message") or "").strip()
            for phrase in ("Confirm which is correct.", "Confirm the temporal relationship."):
                message = message.replace(phrase, "").strip()
            item["question"] = (
                "Some of the details we hold do not agree with each other. "
                f"{message} Please tell us which is correct."
            )
        else:
            message = (check.get("message") or "").strip()
            item["question"] = (
                f"{message} Please provide this information."
                if message
                else "Please provide the missing information for this report."
            )
        items.append(item)

    for number, item in enumerate(items, 1):
        item["number"] = number
    return items


def email_subject(case):
    return f"Follow-up request for adverse reaction report {case['case_number']}"


def email_body(case, product, items, due_date, reminder_of=None, sent_at=None, phone=None):
    product_name = (product or {}).get("product_name") or "a medicine"
    received = _fmt(case.get("received_date"))
    phone = phone or company()["pv_phone"] or DEFAULT_PV_PHONE
    greeting = f"Dear {case['reporter_name']}," if case.get("reporter_name") else "Dear reporter,"
    if reminder_of:
        opening = [
            greeting,
            "",
            f"We wrote to you on {format_sent(reminder_of)} about your report of a "
            f"suspected adverse reaction to {product_name} (our reference "
            f"{case['case_number']}). We have not yet received the information "
            "below and would be grateful for your help:",
        ]
    else:
        opening = [
            greeting,
            "",
            f"Thank you for reporting a suspected adverse reaction to {product_name}"
            + (f", received on {received}" if received else "")
            + f" (our reference {case['case_number']}).",
            "",
            "To complete our assessment we need a little more information:",
        ]
    lines = opening
    lines += [f"  {item['number']}. {item['label']}" for item in items]
    lines += [
        "",
        "The attached form explains each question. Please complete it and "
        "reply to this email"
        + (f" by {_fmt(due_date)}" if due_date else "")
        + ". If some information is not available, please say so.",
        "",
        *([f"If it is easier, you can call us on {phone}.", ""] if phone else []),
        "Kind regards,",
        company()["pv_team_name"],
        company()["letter_name"],
        *([f"Tel: {phone}"] if phone else []),
    ]
    if company()["pv_email"]:
        lines.append(f"Email: {company()['pv_email']}")
    if sent_at:
        lines += ["", f"Sent on {format_sent(sent_at)}."]
    return "\n".join(lines)


# --------------------------------------------------------------------------
# Word form
# --------------------------------------------------------------------------

def _heading(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(APDL_PURPLE)


def _label_table(document, rows):
    table = document.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    table.autofit = False
    for index, (label, value) in enumerate(rows):
        set_cell_shading(table.cell(index, 0), LIGHT_PURPLE)
        set_cell_text(table.cell(index, 0), label, bold=True, size=8, color=APDL_PURPLE)
        set_cell_text(table.cell(index, 1), str(value) if value not in (None, "") else "", size=9)
        table.cell(index, 0).width = Inches(2.1)
        table.cell(index, 1).width = Inches(4.7)
    return table


def _answer_area(document, item):
    if item["answer"] in ("options", "multi"):
        for option in item["options"]:
            line = document.add_paragraph(style=None)
            line.paragraph_format.space_after = Pt(1)
            line.paragraph_format.left_indent = Inches(0.25)
            line.add_run(f"☐  {option}").font.size = Pt(10)
        details = document.add_paragraph()
        details.paragraph_format.space_before = Pt(6)
        details.add_run("Details: ").bold = True
        details.add_run("_" * 70)
    elif item["answer"] == "date":
        line = document.add_paragraph()
        line.add_run("Date (DD/MM/YYYY): ").bold = True
        line.add_run("____ / ____ / ________")
        note = document.add_paragraph()
        note.add_run("If estimated, explain: ").bold = True
        note.add_run("_" * 55)
    else:
        table = document.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        set_cell_text(table.cell(0, 0), "\n\n\n", size=10)


def build_follow_up_request_docx(case, product, items, due_date, prepared_by=None,
                                today=None, sent_at=None, phone=None):
    """The consolidated follow-up form for one case, as a Word file."""
    today = today or date.today()
    phone = phone or company()["pv_phone"] or DEFAULT_PV_PHONE
    product = product or {}
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(10)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run(company()["letter_name"].upper())
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string("666666")
    subtitle = section.header.add_paragraph(company()["department"].upper())
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(9)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("CASE FOLLOW-UP FORM")
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor.from_string(APDL_PURPLE)

    _label_table(
        document,
        form_control()
        + (
            ("SENT ON", format_sent(sent_at)) if sent_at else ("DATE OF REQUEST", _fmt(today)),
            ("PLEASE REPLY BY", _fmt(due_date)),
        ),
    )

    _heading(document, "PART A: THE REPORT THIS REQUEST IS ABOUT")
    patient = case.get("patient_initials") or ""
    if case.get("patient_age_years") is not None:
        patient = (patient + f", {case['patient_age_years']} years").strip(", ")
    if case.get("patient_sex"):
        patient = (patient + f", {case['patient_sex'].lower()}").strip(", ")
    _label_table(
        document,
        (
            ("OUR REFERENCE", case["case_number"]),
            ("REPORT RECEIVED", _fmt(case.get("received_date"))),
            ("REPORTED BY", case.get("reporter_name") or ""),
            ("PATIENT", patient or "Not yet identified"),
            ("SUSPECTED MEDICINE", product.get("product_name") or "Not yet identified"),
            ("REACTION AS REPORTED", case.get("event_description") or "Not yet described"),
        ),
    )

    intro = document.add_paragraph()
    intro.paragraph_format.space_before = Pt(10)
    intro.add_run(
        f"Thank you for your report. To complete our safety assessment we need "
        f"the {len(items)} item(s) below. Please answer what you can; if "
        "something is not known or not available, write \"Not known\"."
    ).font.size = Pt(10)

    _heading(document, "PART B: INFORMATION REQUESTED")
    for item in items:
        question = document.add_paragraph()
        question.paragraph_format.space_before = Pt(10)
        question.paragraph_format.space_after = Pt(4)
        number = question.add_run(f"{item['number']}. {item['label']}")
        number.bold = True
        number.font.color.rgb = RGBColor.from_string(APDL_PURPLE)
        text = document.add_paragraph()
        text.paragraph_format.space_after = Pt(4)
        text.add_run(item["question"]).font.size = Pt(10)
        _answer_area(document, item)

    _heading(document, "PART C: COMPLETED BY")
    _label_table(
        document,
        (
            ("NAME", ""),
            ("PROFESSION", ""),
            ("TELEPHONE / EMAIL", ""),
            ("SIGNATURE AND DATE", ""),
        ),
    )
    back = document.add_paragraph()
    back.paragraph_format.space_before = Pt(8)
    back.add_run(
        "Please return this form by replying to the email it came with, or to "
        f"the {company()['short_name']} Pharmacovigilance team"
        f"{f' (telephone {phone})' if phone else ''}. Information you "
        "provide is used only for medicine safety monitoring."
    ).font.size = Pt(8)

    document.add_page_break()
    _heading(document, "INTERNAL USE: SIGNATORIES")
    signatories = (
        ("Prepared by", (prepared_by or "").upper(), ""),
        ("Reviewed by",) + form_signatories()[0],
        ("Authorised by",) + form_signatories()[1],
    )
    table = document.add_table(rows=len(signatories) + 1, cols=5)
    table.style = "Table Grid"
    for column, label in enumerate(("Signatory", "Name", "Designation", "Signature", "Date")):
        set_cell_shading(table.cell(0, column), LIGHT_PURPLE)
        set_cell_text(table.cell(0, column), label, bold=True, size=8, color=APDL_PURPLE)
    for row, (role, name, designation) in enumerate(signatories, 1):
        set_cell_text(table.cell(row, 0), role, size=8)
        set_cell_text(table.cell(row, 1), name, size=8)
        set_cell_text(table.cell(row, 2), designation, size=8)
        set_cell_text(table.cell(row, 3), "\n\n", size=8)
        set_cell_text(table.cell(row, 4), "\n", size=8)

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output


def request_filename(case):
    return f"follow-up-request-{case['case_number']}.docx"


def group_tasks_by_case(tasks, last_requests=None):
    """Open follow-up tasks grouped into one entry per case, oldest due first.

    Each entry splits the tasks into reporter questions and PV-team tasks.
    """
    from email.utils import parseaddr

    last_requests = last_requests or {}
    groups = {}
    for task in tasks or []:
        group = groups.setdefault(task["case_id"], {
            "case_id": task["case_id"],
            "case_number": task.get("case_number"),
            "workflow_status": task.get("workflow_status"),
            "reporter_name": task.get("reporter_name"),
            "reporter_email": (task.get("reporter_email") or "").strip(),
            "due_date": task.get("due_date"),
            "reporter_tasks": [],
            "check_tasks": [],
            "internal_tasks": [],
            "last_request": last_requests.get(task["case_id"]),
        })
        if task.get("due_date") and (not group["due_date"] or task["due_date"] < group["due_date"]):
            group["due_date"] = task["due_date"]
        task = dict(task)
        task["edit_section"] = EDIT_SECTION_FOR_CODE.get(task.get("check_code"))
        task["edit_field"] = EDIT_FIELD_FOR_CODE.get(task.get("check_code"))
        if task.get("check_code") in INTERNAL_CODES:
            bucket = "internal_tasks"
        elif task.get("check_code") in CHECK_FIRST_CODES:
            bucket = "check_tasks"
        else:
            bucket = "reporter_tasks"
        group[bucket].append(task)
    for group in groups.values():
        email = group["reporter_email"]
        group["can_email"] = bool(email) and parseaddr(email)[1] == email and "@" in email
    return sorted(groups.values(), key=lambda g: (g["due_date"] or date.max, g["case_number"] or ""))
