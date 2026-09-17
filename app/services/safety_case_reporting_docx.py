from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from app.services.ai_case_assessment_docx import (
    APDL_PURPLE,
    LIGHT_PURPLE,
    add_label_value_row,
    set_cell_shading,
    set_cell_text,
)


def _display_date(value):
    if not value:
        return "Not recorded"

    return value.strftime("%d %b %Y")


def _filter_text(report_filters):
    labels = (
        ("Product", report_filters.get("product")),
        ("Country / market", report_filters.get("country")),
        ("Status", report_filters.get("status")),
        ("Priority", report_filters.get("priority")),
    )

    selected = [
        f"{label}: {value}"
        for label, value in labels
        if value
    ]

    return " | ".join(selected) if selected else "No additional filters"


def build_safety_case_reporting_docx(
    cases,
    report_filters,
    editable_content=None,
):
    document = Document()
    editable_content = editable_content or {}
    control_values = editable_content.get("control_values") or []
    summary_values = editable_content.get("summary_values") or []
    section_titles = editable_content.get("section_titles") or []
    edited_case_rows = editable_content.get("case_rows") or []
    edited_signature_rows = (
        editable_content.get("signature_rows") or []
    )

    def saved_value(values, index, fallback):
        if index < len(values) and values[index]:
            return values[index]

        return fallback
    section = document.sections[0]
    report_title = (
        editable_content.get("title")
        or "SAFETY CASE REPORTING SUMMARY"
    )

    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_paragraph.paragraph_format.space_after = Pt(0)

    run = header_paragraph.add_run(
        "ABACUS PARENTERAL DRUGS LIMITED"
    )
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string("666666")

    subtitle = header.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)

    run = subtitle.add_run("REGULATORY AFFAIRS DEPARTMENT")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("666666")

    control_table = document.add_table(rows=4, cols=2)
    control_table.style = "Table Grid"
    control_table.autofit = False

    add_label_value_row(
        control_table,
        0,
        "TITLE",
        saved_value(control_values, 0, report_title),
    )
    add_label_value_row(
        control_table,
        1,
        "REPORTING PERIOD",
        saved_value(
            control_values,
            1,
            report_filters.get("reporting_period") or "All dates",
        ),
    )
    add_label_value_row(
        control_table,
        2,
        "APPLIED FILTERS",
        saved_value(
            control_values,
            2,
            _filter_text(report_filters),
        ),
    )
    add_label_value_row(
        control_table,
        3,
        "REPORT STATUS",
        saved_value(
            control_values,
            3,
            "System-generated filtered report",
        ),
    )

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(16)
    title.paragraph_format.space_after = Pt(12)

    run = title.add_run(report_title)
    run.bold = True
    run.font.size = Pt(15)

    serious_cases = sum(1 for case in cases if case["seriousness"])
    routine_cases = len(cases) - serious_cases

    overview = document.add_table(rows=1, cols=3)
    overview.style = "Table Grid"

    overview_values = (
        (
            "TOTAL SAFETY CASES",
            saved_value(summary_values, 0, len(cases)),
        ),
        (
            "SERIOUS CASES",
            saved_value(summary_values, 1, serious_cases),
        ),
        (
            "ROUTINE CASES",
            saved_value(summary_values, 2, routine_cases),
        ),
    )

    for index, (label_text, value) in enumerate(overview_values):
        cell = overview.cell(0, index)
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)

        label = paragraph.add_run(f"{label_text}\n")
        label.bold = True
        label.font.size = Pt(8)
        label.font.color.rgb = RGBColor.from_string(APDL_PURPLE)

        number = paragraph.add_run(str(value))
        number.bold = True
        number.font.size = Pt(15)

        set_cell_shading(cell, LIGHT_PURPLE)

    document.add_paragraph()

    heading = document.add_paragraph()
    heading.paragraph_format.space_after = Pt(6)

    run = heading.add_run(
        saved_value(
            section_titles,
            0,
            "1. FILTERED SAFETY CASE REGISTER",
        )
    )
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(APDL_PURPLE)

    case_table = document.add_table(rows=1, cols=7)
    case_table.style = "Table Grid"
    case_table.autofit = False

    headings = (
        "Case number",
        "Product",
        "Country",
        "Event",
        "Status",
        "Received",
        "Priority",
    )

    for index, heading_text in enumerate(headings):
        set_cell_shading(case_table.cell(0, index), APDL_PURPLE)
        set_cell_text(
            case_table.cell(0, index),
            heading_text,
            bold=True,
            size=7,
            color="FFFFFF",
        )

    case_rows = edited_case_rows

    if not case_rows:
        case_rows = [
            [
                case["case_number"],
                case["product_name"] or "Not recorded",
                case["country_name"] or "Not recorded",
                case["event_description"] or "Not recorded",
                case["workflow_status"],
                _display_date(case["received_date"]),
                "Serious" if case["seriousness"] else "Routine",
            ]
            for case in cases
        ]

    if case_rows:
        for row in case_rows:
            cells = case_table.add_row().cells

            for index in range(7):
                value = row[index] if index < len(row) else ""
                set_cell_text(cells[index], value, size=7)
    else:
        cells = case_table.add_row().cells
        set_cell_text(
            cells[0],
            "No safety cases match the selected filters",
            size=8,
        )

        for index in range(1, 7):
            set_cell_text(cells[index], "", size=8)

    document.add_paragraph()

    signature_table = document.add_table(rows=4, cols=5)
    signature_table.style = "Table Grid"
    signature_table.autofit = False

    headings = (
        "Signatory",
        "Name",
        "Designation",
        "Signature",
        "Date",
    )

    for index, heading_text in enumerate(headings):
        set_cell_shading(signature_table.cell(0, index), APDL_PURPLE)
        set_cell_text(
            signature_table.cell(0, index),
            heading_text,
            bold=True,
            size=8,
            color="FFFFFF",
        )

    signature_rows = edited_signature_rows or (
        (
            "Prepared by",
            "AMEKO CHARLES",
            "DEPUTY Q.P.P.V.",
            "",
            "",
        ),
        (
            "Reviewed by",
            "YONA ODOI",
            "Q.P.P.V.",
            "",
            "",
        ),
        (
            "Authorised by",
            "KEITH ARUHO",
            "GROUP HEAD, RA & QUALITY",
            "",
            "",
        ),
    )

    for row_index, row in enumerate(signature_rows, start=1):
        values = list(row) + [""] * 5

        for column_index in range(5):
            set_cell_text(
                signature_table.cell(row_index, column_index),
                values[column_index],
                size=8,
            )

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = footer_paragraph.add_run(
        "System-generated safety case reporting summary"
    )
    run.italic = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("666666")

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output