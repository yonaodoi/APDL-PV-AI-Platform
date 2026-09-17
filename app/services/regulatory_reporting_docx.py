from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from app.services.ai_case_assessment_docx import (
    APDL_PURPLE,
    LIGHT_PURPLE,
    add_label_value_row,
    set_cell_shading,
    set_cell_text,
)


def build_regulatory_reporting_docx(
    start_date,
    end_date,
    selected_product,
    metrics,
    monthly_trend,
    product_summary,
):
    document = Document()
    section = document.sections[0]

    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_paragraph.paragraph_format.space_after = Pt(0)

    run = header_paragraph.add_run(
        "ABACUS PARENTERAL DRUGS LIMITED"
    )
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string("666666")

    header_subtitle = header.add_paragraph()
    header_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_subtitle.paragraph_format.space_after = Pt(8)

    run = header_subtitle.add_run(
        "REGULATORY AFFAIRS DEPARTMENT"
    )
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("666666")

    control_table = document.add_table(rows=4, cols=2)
    control_table.style = "Table Grid"
    control_table.autofit = False
    control_table.columns[0].width = Inches(1.6)
    control_table.columns[1].width = Inches(5.8)

    add_label_value_row(
        control_table,
        0,
        "TITLE",
        "REGULATORY REPORTING SUMMARY",
    )
    add_label_value_row(
        control_table,
        1,
        "REPORTING PERIOD",
        f"{start_date:%d %B %Y} to {end_date:%d %B %Y}",
    )
    add_label_value_row(
        control_table,
        2,
        "PRODUCT FILTER",
        selected_product or "All products",
    )
    add_label_value_row(
        control_table,
        3,
        "REPORT STATUS",
        "System-generated management summary",
    )

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(16)
    title.paragraph_format.space_after = Pt(12)

    run = title.add_run("REGULATORY REPORTING SUMMARY")
    run.bold = True
    run.font.size = Pt(15)

    overview = document.add_table(rows=3, cols=2)
    overview.style = "Table Grid"
    overview.autofit = False
    overview.columns[0].width = Inches(3.7)
    overview.columns[1].width = Inches(3.7)

    overview_rows = (
        ("TOTAL RECORDS", metrics["total_reports"]),
        ("SAFETY CASES", metrics["safety_cases"]),
        ("PRODUCT COMPLAINTS", metrics["complaints"]),
        ("SAFETY SIGNALS", metrics["signals"]),
        ("PSURs", metrics["psurs"]),
        (
            "SERIOUS, CRITICAL OR HIGH-PRIORITY",
            metrics["high_priority_reports"],
        ),
    )

    for index, values in enumerate(overview_rows):
        row = index // 2
        column = index % 2
        cell = overview.cell(row, column)
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)

        label = paragraph.add_run(f"{values[0]}\n")
        label.bold = True
        label.font.size = Pt(8)
        label.font.color.rgb = RGBColor.from_string(APDL_PURPLE)

        value = paragraph.add_run(str(values[1] or 0))
        value.bold = True
        value.font.size = Pt(15)

        set_cell_shading(cell, LIGHT_PURPLE)

    document.add_paragraph()

    heading = document.add_paragraph()
    heading.paragraph_format.space_after = Pt(6)
    run = heading.add_run("1. MONTHLY REPORTING TREND")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(APDL_PURPLE)

    monthly_table = document.add_table(rows=1, cols=2)
    monthly_table.style = "Table Grid"

    for index, heading_text in enumerate(
        ("Reporting month", "Total records")
    ):
        set_cell_shading(monthly_table.cell(0, index), APDL_PURPLE)
        set_cell_text(
            monthly_table.cell(0, index),
            heading_text,
            bold=True,
            size=8,
            color="FFFFFF",
        )

    if monthly_trend:
        for row in monthly_trend:
            cells = monthly_table.add_row().cells
            set_cell_text(cells[0], row["period_label"], size=9)
            set_cell_text(cells[1], row["total_reports"], size=9)
    else:
        cells = monthly_table.add_row().cells
        set_cell_text(
            cells[0],
            "No records in selected period",
            size=9,
        )
        set_cell_text(cells[1], "0", size=9)

    document.add_paragraph()

    heading = document.add_paragraph()
    heading.paragraph_format.space_after = Pt(6)
    run = heading.add_run("2. PRODUCT-WISE REPORTING SUMMARY")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(APDL_PURPLE)

    summary_table = document.add_table(rows=1, cols=7)
    summary_table.style = "Table Grid"
    summary_table.autofit = False

    headers = (
        "Product",
        "Total",
        "Safety cases",
        "Complaints",
        "Signals",
        "PSURs",
        "Priority",
    )

    for index, heading_text in enumerate(headers):
        set_cell_shading(summary_table.cell(0, index), APDL_PURPLE)
        set_cell_text(
            summary_table.cell(0, index),
            heading_text,
            bold=True,
            size=7,
            color="FFFFFF",
        )

    if product_summary:
        for row in product_summary:
            cells = summary_table.add_row().cells
            set_cell_text(cells[0], row["product_name"], size=7)
            set_cell_text(cells[1], row["total_reports"], size=7)
            set_cell_text(cells[2], row["safety_cases"], size=7)
            set_cell_text(cells[3], row["complaints"], size=7)
            set_cell_text(cells[4], row["signals"], size=7)
            set_cell_text(cells[5], row["psurs"], size=7)
            set_cell_text(
                cells[6],
                row["high_priority_reports"],
                size=7,
            )
    else:
        cells = summary_table.add_row().cells
        set_cell_text(
            cells[0],
            "No records in selected period",
            size=7,
        )
        for index in range(1, 7):
            set_cell_text(cells[index], "0", size=7)

    document.add_paragraph()

    signature_table = document.add_table(rows=4, cols=5)
    signature_table.style = "Table Grid"
    signature_table.autofit = False

    column_widths = (
        Inches(1.05),
        Inches(1.2),
        Inches(2.25),
        Inches(1.2),
        Inches(1.15),
    )

    for index, width in enumerate(column_widths):
        for cell in signature_table.columns[index].cells:
            cell.width = width

    for index, heading_text in enumerate(
        ("Signatory", "Name", "Designation", "Signature", "Date")
    ):
        set_cell_shading(signature_table.cell(0, index), APDL_PURPLE)
        set_cell_text(
            signature_table.cell(0, index),
            heading_text,
            bold=True,
            size=8,
            color="FFFFFF",
        )

    signature_rows = (
        ("Prepared by", "AMEKO CHARLES", "DEPUTY Q.P.P.V."),
        ("Reviewed by", "YONA ODOI", "Q.P.P.V."),
        (
            "Authorised by",
            "KEITH ARUHO",
            "GROUP HEAD, RA & QUALITY",
        ),
    )

    for row_index, values in enumerate(signature_rows, start=1):
        set_cell_text(signature_table.cell(row_index, 0), values[0], size=8)
        set_cell_text(signature_table.cell(row_index, 1), values[1], size=8)
        set_cell_text(signature_table.cell(row_index, 2), values[2], size=8)
        set_cell_text(signature_table.cell(row_index, 3), "", size=8)
        set_cell_text(signature_table.cell(row_index, 4), "", size=8)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = footer_paragraph.add_run(
        "System-generated regulatory reporting summary"
    )
    run.italic = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("666666")

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output