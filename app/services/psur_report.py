from datetime import date
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph
from flask import current_app
from app.db import query_all
from app.psur.section_definitions import PSUR_SECTION_TITLES
DEFAULT_TEMPLATE_SECTION_HEADINGS = (
    "Executive summary",
    "Introduction",
    "2.0 Use in special population:",
    "3.0 Worldwide marketing authorisation status",
    "4.0 Action taken in the reporting interval for safety reasons",
    "5.0 Changes to reference safety information",
    "6.0 Estimated Exposure and Use Patterns",
    "6.1 Cumulative Subject Exposure in Clinical Trials",
    "6.2 Cumulative and Interval Patient Exposure from Marketing Experience",
    "6.2.1. Post-approval (non-clinical trial) exposure",
    "6.2.2. Post-Authorisation uses in special populations",
    "6.2.3. Other post-authorisation use",
    "Off-label use:",
    "Overdose:",
    "Drug abuse and misuse:",
    "Data in Summary Tabulations",
    "7.1 Reference Information",
    "7.2 Cumulative Summary Tabulations of Serious Adverse Events from Clinical Trials",
    "7.3 Cumulative and Interval Summary Tabulations form Post-Marketing Data Sources",
    "Summaries of significant findings from Clinical Trials in the reporting interval",
    "8.1 Completed clinical trials",
    "8.2 Ongoing clinical trials",
    "8.3 Long-term follow-up",
    "8.4 Other Therapeutic use of medicinal product",
    "8.5 New safety data related to fixed combination therapies",
    "9.0 Findings from non-interventional studies",
    "10. Information from other clinical trials and sources",
    "10.1 Other clinical trials",
    "10.2 Medication error",
    "10.3 Analysis of other events",
    "11. Non-Clinical Data",
    "12. Literature",
    "12.1 Literature publications on lack of efficacy",
    "12.2 Literature publications on overdose, abuse or misuse",
    "12.3 Literature publications with compassionate supply, named patient use",
    "12.4 Literature publications with medication error where no adverse events occurred",
    "12.5 Literature publications on pregnancy outcomes (including termination) with/without adverse outcomes",
    "12.6 Literature publications with use in paediatric populations",
    "12.7 Literature publications with important non-clinical safety results",
    "12.8 Other relevant literature publications",
    "13.0 Other periodic reports",
    "14.0 Lack of efficacy in controlled Clinical Trials",
    "15.0 Late breaking information",
    "16.0 Overview of signals: New, ongoing or closed",
    "17. Signal and risk evaluation",
    "17.1 Summary of safety concerns",
    "17.2 Signal evaluation",
    "17.3 Evaluation of risks and new information",
    "17.4 Characterisation of risks",
    "17.5 Effectiveness of Risk Minimisation",
    "18.0 Benefit evaluation",
    "18.1 Important baseline efficacy and effectiveness information",
    "18.2 Newly identified information on efficacy and effectiveness",
    "18.3 Characterisation of benefits",
    "19. Integrated benefit-risk analysis for authorised indications",
    "19.1 Benefit-risk context – medical need and important alternatives",
    "19.2 Benefit-Risk analysis evaluation",
    "20.0 Conclusion and actions",
    "21.0 Appendices",
)

SECTION_HEADINGS = PSUR_SECTION_TITLES


def _text(value, fallback="Not recorded"):
    if value is None:
        return fallback

    value = str(value).strip()
    return value or fallback


def _date(value):
    if not value:
        return "Not recorded"

    if isinstance(value, date):
        return value.strftime("%d %B %Y")

    return str(value)


def _replace_in_paragraph(paragraph, old_text, new_text):
    for run in paragraph.runs:
        if old_text in run.text:
            run.text = run.text.replace(old_text, new_text)
            return True

    if old_text in paragraph.text:
        paragraph.text = paragraph.text.replace(old_text, new_text)
        return True

    return False


def _replace_paragraph_start(paragraph, label, value):
    if paragraph.text.strip().startswith(label):
        _replace_in_paragraph(
            paragraph,
            paragraph.text,
            f"{label} {value}",
        )
        return True

    return False


def _insert_paragraph_after(paragraph, text):
    new_paragraph_xml = OxmlElement("w:p")
    paragraph._p.addnext(new_paragraph_xml)

    new_paragraph = Paragraph(
        new_paragraph_xml,
        paragraph._parent,
    )
    new_paragraph.add_run(text)

    return new_paragraph


def _safe_filename(value):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", value)
def _set_cell_shading(cell, fill):
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell_properties.append(shading)


def _set_cell_text(cell, value, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    run = paragraph.add_run(str(value or ""))
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(10)

    if color:
        run.font.color.rgb = RGBColor.from_string(color)

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _append_signatory_table(
    document,
    prepared_by,
    qppv_name,
    group_head_name,
):
    document.add_paragraph()

    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(4)
    title.paragraph_format.space_after = Pt(6)

    run = title.add_run("SIGNATORIES")
    run.bold = True
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)

    table = document.add_table(rows=4, cols=5)
    table.style = "Table Grid"
    table.autofit = False

    headings = (
        "Signatory",
        "Name",
        "Designation",
        "Signature",
        "Date",
    )

    for index, heading in enumerate(headings):
        cell = table.cell(0, index)
        _set_cell_shading(cell, "5F3ED4")
        _set_cell_text(
            cell,
            heading,
            bold=True,
            color="FFFFFF",
        )

    signatories = (
        (
            "Prepared by",
            prepared_by,
            "DEPUTY Q.P.P.V.",
        ),
        (
            "Reviewed by",
            qppv_name,
            "Q.P.P.V.",
        ),
        (
            "Authorised by",
            group_head_name,
            "GROUP HEAD, RA & QUALITY",
        ),
    )

    for row_index, signatory in enumerate(signatories, start=1):
        _set_cell_text(table.cell(row_index, 0), signatory[0])
        _set_cell_text(table.cell(row_index, 1), signatory[1])
        _set_cell_text(table.cell(row_index, 2), signatory[2])
        _set_cell_text(table.cell(row_index, 3), "")
        _set_cell_text(table.cell(row_index, 4), "")


def _format_report_document(document):
    headings = set(DEFAULT_TEMPLATE_SECTION_HEADINGS)
    headings.update(SECTION_HEADINGS.values())
    headings.update(
        {
            "Therapeutic Indication:",
            "Mechanism of action:",
            "SIGNATORIES",
        }
    )

    for paragraph in document.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0

        is_heading = paragraph.text.strip() in headings

        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(
                qn("w:eastAsia"),
                "Times New Roman",
            )
            run.font.size = Pt(12)

            if is_heading:
                run.bold = True

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1.0

                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run._element.rPr.rFonts.set(
                            qn("w:eastAsia"),
                            "Times New Roman",
                        )
                        run.font.size = Pt(10)

    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            paragraph.text = ""
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)

        for table in list(section.footer.tables):
            table._element.getparent().remove(table._element)

def _section_default_content(section_key, report):
    product_name = _text(report["product_name"])
    period = (
        f"{_date(report['reporting_period_start'])} to "
        f"{_date(report['reporting_period_end'])}"
    )

    if section_key == "special_populations":
        return (
            f"The APDL PV records for {product_name} were reviewed "
            f"for the period {period}. No completed analysis by "
            "special population was available in the platform. "
            "Individual ADR cases should be reviewed for pregnancy, "
            "paediatric use, elderly patients and other clinically "
            "relevant populations where this information is available."
        )

    if section_key in (
        "exposure",
        "clinical_trial_exposure",
        "marketing_exposure",
        "post_approval_exposure",
        "special_population_exposure",
        "other_post_authorisation_use",
    ):
        return (
            "No validated sales, distribution, prescription or "
            "patient-exposure denominator was available in the APDL "
            "PV platform for this reporting interval. Reporting "
            "rates and exposure-adjusted comparisons have therefore "
            "not been calculated. Verified exposure data should be "
            "obtained from the relevant commercial or supply records."
        )

    if section_key in (
        "off_label_use",
        "overdose",
        "abuse_and_misuse",
        "medication_error",
    ):
        return (
            "The APDL PV records did not contain a completed, "
            "section-specific review for this topic during the "
            "reporting interval. The QPPV should confirm the position "
            "from case narratives, complaint investigations, medical "
            "information and any other relevant source before "
            "finalising the PBRER."
        )

    if section_key in (
        "clinical_trial_serious_events",
        "clinical_trial_findings",
        "completed_clinical_trials",
        "ongoing_clinical_trials",
        "long_term_follow_up",
        "other_therapeutic_use",
        "combination_therapy_safety",
        "other_trials_and_sources",
        "other_clinical_trials",
        "lack_of_efficacy",
    ):
        return (
            "No APDL-sponsored clinical-trial information relevant to "
            "this section was available in the PV platform for the "
            "reporting interval. This section should be completed "
            "from verified clinical-development records where such "
            "studies apply."
        )

    if section_key == "non_interventional_studies":
        return (
            "No findings from non-interventional studies were entered "
            "in the APDL PV platform for this reporting interval. The "
            "QPPV should confirm whether any post-authorisation safety "
            "study, observational study or other real-world evidence "
            "source applies to this product."
        )

    if section_key in (
        "non_clinical_data",
        "literature",
        "literature_lack_of_efficacy",
        "literature_overdose_abuse",
        "literature_compassionate_supply",
        "literature_medication_error",
        "literature_pregnancy",
        "literature_paediatric",
        "literature_non_clinical",
        "other_relevant_literature",
    ):
        return (
            "No verified information for this section was available "
            "from the APDL PV platform during the reporting interval. "
            "Relevant literature-screening, non-clinical or medical "
            "information records should be reviewed and entered here "
            "before the PBRER is finalised."
        )

    if section_key in (
        "other_periodic_reports",
        "late_breaking_information",
        "appendices",
    ):
        return (
            "No information applicable to this section was entered in "
            "the APDL PV platform for the reporting interval. The "
            "QPPV should confirm this against the applicable "
            "regulatory, safety and quality records before approval."
        )

    if section_key in (
        "signals_overview",
        "safety_concerns",
        "signal_evaluation",
        "risk_evaluation",
        "risk_characterisation",
        "risk_minimisation_effectiveness",
    ):
        return (
            "The safety signal and risk-evaluation position should be "
            "determined from the integrated review of ADR cases, "
            "signal-screening results, market complaints, reference "
            "safety information and any external evidence. Complete "
            "the evaluation and record the medical/QPPV conclusion."
        )

    if section_key in (
        "benefit_evaluation",
        "baseline_efficacy",
        "new_efficacy_information",
        "benefit_characterisation",
        "benefit_risk_context",
        "benefit_risk_evaluation",
    ):
        return (
            "No product-specific efficacy or effectiveness evidence "
            "for this section was recorded in the APDL PV platform "
            "during the reporting interval. The benefit assessment "
            "should be completed from the approved product "
            "information, relevant clinical evidence and current "
            "medical knowledge."
        )

    return (
        "Review of the APDL pharmacovigilance records did not identify "
        "information applicable to this section during the reporting "
        "interval. The QPPV should confirm this against other relevant "
        "data sources before the PBRER is finalised."
    )
def generate_psur_report(report):
    template_path = Path(current_app.config["PSUR_TEMPLATE_PATH"])

    if not template_path.exists():
        raise FileNotFoundError(
            f"PSUR template was not found: {template_path}"
        )

    document = Document(template_path)

    replacements = {
        "ACTIVE SUBSTANCE(S): WRITE": (
            "ACTIVE SUBSTANCE(S): "
            + _text(report["active_substances"])
        ),
        "ATC CODE(S): WRITE": (
            "ATC CODE(S): "
            + _text(report["atc_codes"])
        ),
        "MARKETING AUTHORISATION PROCEDURE in the EU: WRITE": (
            "MARKETING AUTHORISATION PROCEDURE in the EU: "
            + _text(report["marketing_authorisation_procedure"])
        ),
        "INTERNATIONAL BIRTH DATE (IBD): WRITE": (
            "INTERNATIONAL BIRTH DATE (IBD): "
            + _date(report["international_birth_date"])
        ),
        "EUROPEAN UNION REFERENCE DATE (EURD): WRITE": (
            "EUROPEAN UNION REFERENCE DATE (EURD): "
            + _date(report["eurd"])
        ),
        "NAME: WRITE": (
            "NAME: " + _text(report["qppv_name"])
        ),
        "Phone: WRITE": (
            "Phone: " + _text(report["qppv_phone"])
        ),
        "Email: write": (
            "Email: " + _text(report["qppv_email"])
        ),
        "QPPV: WRITE NAME": (
            "QPPV: " + _text(report["qppv_name"])
        ),
    }
    qppv_name = _text(
        report["qppv_name"],
        "YONA ODOI",
    )
    group_head_name = _text(
        report["approved_by"],
        "KEITH ARUHO",
    )
    for paragraph in document.paragraphs:
        for old_text, new_text in replacements.items():
            _replace_in_paragraph(paragraph, old_text, new_text)

        _replace_paragraph_start(
            paragraph,
            "MEDICINAL PRODUCTS COVERED:",
            _text(report["product_name"]),
        )

    contact_person_section = False
    reviewer_section = False

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("QPPV"):
            paragraph.text = f"QPPV: {qppv_name}"
            continue

        if text.startswith("Contact person for the PBRER:"):
            contact_person_section = True
            reviewer_section = False
            continue

        if text.startswith("REVIEWER (A):"):
            reviewer_section = True
            contact_person_section = False
            continue

        if contact_person_section and text == "Name:":
            paragraph.text = f"Name: {qppv_name}"
            continue

        if contact_person_section and text == "Position:":
            paragraph.text = "Position: QPPV"
            contact_person_section = False
            continue

        if reviewer_section and text == "Name:":
            paragraph.text = f"Name: {group_head_name}"
            continue

        if reviewer_section and text == "Position:":
            paragraph.text = "Position: GROUP HEAD, RA & QUALITY"
            reviewer_section = False
    introduction_seen = False

    for paragraph in list(document.paragraphs):
        heading = paragraph.text.strip()

        if heading == "Introduction":
            introduction_seen = True
            continue

        if (
            introduction_seen
            and heading in (
                "Therapeutic Indication:",
                "Mechanism of action",
            )
        ):
            paragraph._element.getparent().remove(
                paragraph._element
            )
    therapeutic_indication_added = False
    mechanism_of_action_added = False

    for paragraph in list(document.paragraphs):
        heading = paragraph.text.strip()

        if (
            heading == "Therapeutic Indication:"
            and not therapeutic_indication_added
        ):
            _insert_paragraph_after(
                paragraph,
                _text(report["therapeutic_indication"]),
            )
            therapeutic_indication_added = True

        if (
            heading == "Mechanism of action:"
            and not mechanism_of_action_added
        ):
            _insert_paragraph_after(
                paragraph,
                _text(report["mechanism_of_action"]),
            )
            mechanism_of_action_added = True
    if len(document.tables) >= 1:
        authorisation_table = document.tables[0]

        if len(authorisation_table.rows) > 1:
            first_country = _text(
                report["countries_covered"]
            ).split(",")[0].strip()

            authorisation_table.cell(1, 0).text = _text(
                report["product_name"]
            )
            authorisation_table.cell(1, 1).text = _text(
                report["marketing_authorisation_number"]
            )
            authorisation_table.cell(1, 2).text = first_country
            authorisation_table.cell(1, 3).text = _date(
                report["marketing_authorisation_date"]
            )
            authorisation_table.cell(1, 4).text = _text(
                report["marketing_authorisation_holder_name"]
            )

    if len(document.tables) >= 2:
        period_table = document.tables[1]
        period_text = (
            "PERIOD COVERED BY THIS REPORT "
            f"From {_date(report['reporting_period_start'])} "
            f"to {_date(report['reporting_period_end'])}  "
            f"DATE OF THIS REPORT {_date(date.today())}"
        )
        period_table.cell(0, 0).text = period_text

    if len(document.tables) >= 3:
        serial_table = document.tables[2]

        if len(serial_table.rows) > 1:
            serial_table.cell(1, 0).text = _text(
                report["serial_number"],
                report["report_number"],
            )
            serial_table.cell(1, 1).text = (
                f"{_date(report['reporting_period_start'])} "
                f"to {_date(report['reporting_period_end'])}"
            )

    if len(document.tables) >= 4:
        country_table = document.tables[3]
        countries = _text(report["countries_covered"]).split(",")

        for index, country in enumerate(countries[:2], start=1):
            if index < len(country_table.rows):
                country_table.cell(index, 0).text = _text(
                    report["active_substances"]
                )
                country_table.cell(index, 1).text = country.strip()
                country_table.cell(index, 2).text = _date(
                    report["marketing_authorisation_date"]
                )

    saved_sections = query_all(
        """
        SELECT section_key, content
        FROM pv.psur_section_entries
        WHERE psur_id = %s
        """,
        (report["psur_id"],),
    )

    content_by_key = {
        section["section_key"]: section["content"]
        for section in saved_sections
    }

    for paragraph in list(document.paragraphs):
        heading = paragraph.text.strip()
        section_key = None

        for key, template_heading in SECTION_HEADINGS.items():
            if heading == template_heading:
                section_key = key
                break

        content = content_by_key.get(section_key)

        if content:
            _insert_paragraph_after(paragraph, content)
        elif section_key:
            _insert_paragraph_after(
                paragraph,
                _section_default_content(section_key, report),
            )
        elif heading in DEFAULT_TEMPLATE_SECTION_HEADINGS:
            _insert_paragraph_after(
                paragraph,
                _section_default_content(None, report),
            )

    output_folder = (
        Path(current_app.config["UPLOAD_ROOT"]).resolve()
        / "reports"
    )
    output_folder.mkdir(parents=True, exist_ok=True)

    filename = (
        f"PSUR_{_safe_filename(report['report_number'])}.docx"
    )
    output_path = output_folder / filename
    _append_signatory_table(
        document,
        prepared_by=_text(
            report["prepared_by"],
            "CHARLES AMEKO",
        ),
        qppv_name=qppv_name,
        group_head_name=group_head_name,
    )

    _format_report_document(document)

    document.save(output_path)

    return output_path