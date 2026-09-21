from datetime import date
from io import BytesIO
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from flask import current_app


def _text(value, fallback=""):
    if value is None:
        return fallback

    value = str(value).strip()
    return value or fallback


def _date(value):
    if not value:
        return ""

    if isinstance(value, date):
        return value.strftime("%d %B %Y")

    return str(value)


def _normalise_heading(value):
    value = _text(value).lower()
    value = value.replace(":", "")
    value = re.sub(r"\s+", " ", value)
    return value.strip()


def _set_run_font(run, size=12, bold=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "Times New Roman",
    )
    run.font.size = Pt(size)

    if bold is not None:
        run.bold = bold


def _set_cell_text(
    cell,
    value,
    size=10,
    bold=False,
    color=None,
):
    cell.text = ""

    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)

    run = paragraph.add_run(_text(value))
    _set_run_font(run, size=size, bold=bold)

    if color:
        run.font.color.rgb = color
def _set_cell_shading(cell, fill):
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell_properties.append(shading)
def _append_signatory_table(document, report):
    prepared_by_name = _text(report.get("prepared_by"))
    prepared_by_position = _text(
        report.get("prepared_by_position")
    )
    reviewer_name = _text(report.get("reviewer_a_name"))
    reviewer_position = _text(
        report.get("reviewer_a_position")
    )
    qppv_name = _text(report.get("qppv_name"))

    signatory_heading = None

    for paragraph in document.paragraphs:
        if _normalise_heading(paragraph.text) == "signatories":
            signatory_heading = paragraph
            break

    table = document.add_table(rows=4, cols=5)
    table.style = "Table Grid"

    headings = (
        "Signatory",
        "Name",
        "Position",
        "Signature",
        "Date",
    )

    for index, heading in enumerate(headings):
        cell = table.cell(0, index)
        _set_cell_shading(cell, "5F3ED4")
        _set_cell_text(
            cell,
            heading,
            size=9,
            bold=True,
        )

    signatories = (
        (
            "Prepared by",
            prepared_by_name,
            prepared_by_position,
        ),
        (
            "Reviewed by",
            reviewer_name,
            reviewer_position,
        ),
        (
            "Approved by",
            qppv_name,
            "QPPV",
        ),
    )

    for row_index, signatory in enumerate(signatories, start=1):
        _set_cell_text(table.cell(row_index, 0), signatory[0])
        _set_cell_text(table.cell(row_index, 1), signatory[1])
        _set_cell_text(table.cell(row_index, 2), signatory[2])
        _set_cell_text(table.cell(row_index, 3), "")
        _set_cell_text(table.cell(row_index, 4), "")

    if signatory_heading is not None:
        signatory_heading._p.addnext(table._tbl)

def _replace_label(paragraph, label, value):
    if not paragraph.text.strip().startswith(label):
        return False

    paragraph.text = ""

    run = paragraph.add_run(
        f"{label} {_text(value, 'Not recorded')}"
    )
    _set_run_font(run, size=10)

    return True


def _insert_paragraph_after(paragraph, text):
    new_paragraph_xml = OxmlElement("w:p")
    paragraph._p.addnext(new_paragraph_xml)

    new_paragraph = Paragraph(
        new_paragraph_xml,
        paragraph._parent,
    )
    new_paragraph.paragraph_format.space_before = Pt(0)
    new_paragraph.paragraph_format.space_after = Pt(5)
    new_paragraph.paragraph_format.line_spacing = 1.0

    run = new_paragraph.add_run(text)
    _set_run_font(run, size=12)

    return new_paragraph


def _insert_section_content_after(paragraph, content):
    anchor = paragraph

    for text in _text(content).splitlines():
        anchor = _insert_paragraph_after(anchor, text)

def _fill_front_matter(document, report):
    replacements = {
        "ACTIVE SUBSTANCE(S):": report.get("active_substances"),
        "ATC CODE(S):": report.get("atc_codes"),
        "MEDICINAL PRODUCTS COVERED:": report.get("product_name"),
        "MARKETING AUTHORISATION PROCEDURE in the EU:": (
            report.get("marketing_authorisation_procedure")
        ),
        "INTERNATIONAL BIRTH DATE (IBD):": _date(
            report.get("international_birth_date")
        ),
        "EUROPEAN UNION REFERENCE DATE (EURD):": _date(
            report.get("eurd")
        ),
    }

    qppv_name = _text(report.get("qppv_name"))
    qppv_phone = _text(report.get("qppv_phone"))
    qppv_email = _text(report.get("qppv_email"))

    pbrer_contact_name = _text(
        report.get("pbrer_contact_name"),
        qppv_name,
    )
    pbrer_contact_position = _text(
        report.get("pbrer_contact_position"),
        "QPPV",
    )
    reviewer_name = _text(report.get("reviewer_a_name"))
    reviewer_position = _text(
        report.get("reviewer_a_position")
    )

    qppv_contact_section = False
    contact_person_section = False
    reviewer_section = False
    therapeutic_indication_added = False
    mechanism_of_action_added = False

    for paragraph in document.paragraphs:
        for label, value in replacements.items():
            if _replace_label(paragraph, label, value):
                break

        text = paragraph.text.strip()

        if text.startswith(
            "NAME AND CONTACT DETAILS OF THE QUALIFIED PERSON"
        ):
            qppv_contact_section = True
            continue

        if qppv_contact_section and text.startswith("NAME:"):
            paragraph.text = ""
            run = paragraph.add_run(f"NAME: {qppv_name}")
            _set_run_font(run, size=10)
            continue

        if qppv_contact_section and text.startswith("Phone:"):
            paragraph.text = ""
            run = paragraph.add_run(f"Phone: {qppv_phone}")
            _set_run_font(run, size=10)
            continue

        if qppv_contact_section and text.startswith("Email:"):
            paragraph.text = ""
            run = paragraph.add_run(f"Email: {qppv_email}")
            _set_run_font(run, size=10)
            qppv_contact_section = False
            continue

        if (
            text.startswith("Therapeutic Indication:")
            and not therapeutic_indication_added
        ):
            paragraph.text = ""
            run = paragraph.add_run(
                "Therapeutic Indication: "
                + _text(report.get("therapeutic_indication"))
            )
            _set_run_font(run, size=10)
            therapeutic_indication_added = True
            continue

        if (
            text.startswith("Mechanism of action:")
            and not mechanism_of_action_added
        ):
            paragraph.text = ""
            run = paragraph.add_run(
                "Mechanism of action: "
                + _text(report.get("mechanism_of_action"))
            )
            _set_run_font(run, size=10)
            mechanism_of_action_added = True
            continue

        if text.startswith("QPPV"):
            paragraph.text = ""
            run = paragraph.add_run(f"QPPV Name: {qppv_name}")
            _set_run_font(run, size=10)
            continue

        if text.startswith("Contact person for the PBRER:"):
            contact_person_section = True
            reviewer_section = False
            continue

        if text.startswith("REVIEWER (A):"):
            reviewer_section = True
            contact_person_section = False
            continue

        if contact_person_section and text == "Name:":
            paragraph.text = ""
            run = paragraph.add_run(
                f"Name: {pbrer_contact_name}"
            )
            _set_run_font(run, size=10)
            continue

        if contact_person_section and text == "Position:":
            paragraph.text = ""
            run = paragraph.add_run(
                f"Position: {pbrer_contact_position}"
            )
            _set_run_font(run, size=10)
            contact_person_section = False
            continue

        if reviewer_section and text == "Name:":
            paragraph.text = ""
            run = paragraph.add_run(f"Name: {reviewer_name}")
            _set_run_font(run, size=10)
            continue

        if reviewer_section and text == "Position:":
            paragraph.text = ""
            run = paragraph.add_run(
                f"Position: {reviewer_position}"
            )
            _set_run_font(run, size=10)
            reviewer_section = False

    if len(document.tables) >= 1:
        table = document.tables[0]

        if len(table.rows) > 1:
            country = _text(
                report.get("countries_covered")
            ).split(",")[0].strip()

            _set_cell_text(table.cell(1, 0), report.get("product_name"))
            _set_cell_text(
                table.cell(1, 1),
                report.get("marketing_authorisation_number"),
            )
            _set_cell_text(table.cell(1, 2), country)
            _set_cell_text(
                table.cell(1, 3),
                _date(report.get("marketing_authorisation_date")),
            )
            _set_cell_text(
                table.cell(1, 4),
                report.get("marketing_authorisation_holder_name"),
            )

    if len(document.tables) >= 2:
        _set_cell_text(
            document.tables[1].cell(0, 0),
            (
                "PERIOD COVERED BY THIS REPORT\n"
                f"From: {_date(report.get('reporting_period_start'))} "
                f"to: {_date(report.get('reporting_period_end'))}\n"
                f"DATE OF THIS REPORT: {_date(date.today())}"
            ),
        )

    if len(document.tables) >= 3:
        table = document.tables[2]

        if len(table.rows) > 1:
            _set_cell_text(
                table.cell(1, 0),
                report.get("serial_number")
                or report.get("report_number"),
            )
            _set_cell_text(
                table.cell(1, 1),
                (
                    f"{_date(report.get('reporting_period_start'))} "
                    f"to {_date(report.get('reporting_period_end'))}"
                ),
            )

    if len(document.tables) >= 4:
        table = document.tables[3]

        if len(table.rows) > 1:
            country = _text(
                report.get("countries_covered")
            ).split(",")[0].strip()

            _set_cell_text(
                table.cell(1, 0),
                (
                    f"{_text(report.get('active_substances'))} / "
                    f"{_text(report.get('product_name'))}"
                ),
            )
            _set_cell_text(table.cell(1, 1), country)
            _set_cell_text(
                table.cell(1, 2),
                _date(report.get("marketing_authorisation_date")),
            )

def build_psur_builder_docx(report, sections):
    template_path = current_app.config["PSUR_TEMPLATE_PATH"]

    if not template_path.exists():
        raise FileNotFoundError(
            f"PSUR template was not found: {template_path}"
        )

    document = Document(template_path)
    _fill_front_matter(document, report)

    content_by_heading = {
        _normalise_heading(section["section_title"]): (
            section["final_content"]
        )
        for section in sections
        if section["final_content"]
    }

    for paragraph in list(document.paragraphs):
        heading = _normalise_heading(paragraph.text)

        if heading in content_by_heading:
            _insert_section_content_after(
                paragraph,
                content_by_heading[heading],
            )
    _append_signatory_table(document, report)

    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            paragraph.text = ""

        for table in list(section.footer.tables):
            table._element.getparent().remove(table._element)

    output = BytesIO()
    document.save(output)
    output.seek(0)

    return output