from pathlib import Path

from docx import Document
from flask import current_app


def _date(value):
    if not value:
        return "___ / ___ / _____"
    return value.strftime("%d / %m / %Y")


def _time(value):
    if not value:
        return "Not recorded"
    return value.strftime("%H:%M")


def _value(value):
    return str(value).strip() if value else "Not recorded"


def _checked(label, selected):
    mark = "☒" if selected else "☐"
    return f"{mark} {label}"


def _find_paragraph(document, text_start):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(text_start):
            return paragraph
    return None


def _replace_after(document, heading_start, value):
    paragraphs = document.paragraphs

    for index, paragraph in enumerate(paragraphs):
        if paragraph.text.strip().startswith(heading_start):
            if index + 1 < len(paragraphs):
                paragraphs[index + 1].text = value
            return


def _contains(value, phrase):
    return phrase.lower() in (value or "").lower()


def generate_adr_report(case, product):
    template_path = Path(current_app.config["ADR_TEMPLATE_PATH"])

    if not template_path.exists():
        raise FileNotFoundError(
            f"ADR template was not found: {template_path}"
        )

    document = Document(template_path)

    report_type = case.get("report_type") or ""
    serious = bool(case.get("seriousness"))

    type_paragraph = _find_paragraph(document, "☐ Initial")
    if type_paragraph:
        type_paragraph.text = "     ".join(
            [
                _checked("Initial", report_type == "Initial"),
                _checked("Follow up", report_type == "Follow-up"),
                _checked("Serious", serious),
                _checked("Not Serious", not serious),
                _checked("Drug", True),
            ]
        )

    patient_paragraph = _find_paragraph(document, "Patient ID/Initials:")
    if patient_paragraph:
        patient_paragraph.text = (
            f"Patient ID/Initials: {_value(case.get('patient_initials'))}     "
            f"Gender: {_checked('Male', case.get('patient_sex') == 'Male')}   "
            f"{_checked('Female', case.get('patient_sex') == 'Female')}     "
            f"Weight (kg): {_value(case.get('patient_weight_kg'))}"
        )

    pregnancy_status = case.get("patient_pregnancy_status") or ""
    _replace_after(
        document,
        "Pregnancy Status:",
        "   ".join(
            [
                _checked("Yes", pregnancy_status == "Yes"),
                _checked("No", pregnancy_status == "No"),
                _checked(
                    "N/A",
                    pregnancy_status in ("Not applicable", "Unknown", ""),
                ),
            ]
        ),
    )

    address_paragraph = _find_paragraph(document, "Full Address:")
    if address_paragraph:
        address_paragraph.text = (
            f"Full Address: {_value(case.get('patient_address'))}     "
            f"Telephone Number: {_value(case.get('patient_phone'))}"
        )

    birth_paragraph = _find_paragraph(document, "Date of Birth:")
    if birth_paragraph:
        birth_paragraph.text = (
            f"Date of Birth: {_date(case.get('patient_date_of_birth'))}   OR   "
            f"Age at onset: {_value(case.get('patient_age_years'))}"
        )

    _replace_after(
        document,
        "Medical History:",
        _value(case.get("medical_history")),
    )

    if document.tables:
        table = document.tables[0]
        if len(table.rows) > 1:
            row = table.rows[1]
            values = [
                _value(product.get("generic_name")),
                _value(product.get("product_name")),
                _value(product.get("batch_number")),
                " | ".join(
                    filter(
                        None,
                        [
                            product.get("route"),
                            product.get("dose"),
                            product.get("frequency"),
                            f"Strength: {product.get('strength')}" if product.get("strength") else None,
                        ],
                    )
                )
                or "Not recorded",
                _date(product.get("therapy_start_date")),
                _date(product.get("therapy_end_date")),
                _value(product.get("indication")),
                "☒",
            ]

            for cell, value in zip(row.cells, values):
                cell.text = value

    _replace_after(
        document,
        "4.0 Brief Description",
        (
            f"ADR: {_value(case.get('event_description'))}\n\n"
            f"Treatment given: {_value(case.get('treatment_given'))}"
        ),
    )

    onset_paragraph = _find_paragraph(document, "Date of ADR onset:")
    if onset_paragraph:
        onset_paragraph.text = (
            f"Date of ADR onset: {_date(case.get('event_onset_date'))}   "
            f"Time of onset: {_time(case.get('event_onset_time'))}   "
            f"Date ended: {_date(case.get('event_end_date'))}"
        )

    _replace_after(
        document,
        "5.0 Relevant Laboratory Test Results",
        f"{_value(case.get('laboratory_results'))}\n",
    )

    seriousness_text = case.get("seriousness_criteria") or ""
    serious_paragraph = _find_paragraph(document, "☐ Prolonged hospitalization")
    if serious_paragraph:
        serious_paragraph.text = "     ".join(
            [
                _checked(
                    "Prolonged hospitalization",
                    _contains(seriousness_text, "hospital"),
                ),
                _checked(
                    "Disability",
                    _contains(seriousness_text, "disability"),
                ),
                _checked("Death", _contains(seriousness_text, "death")),
                _checked(
                    "Life-threatening",
                    _contains(seriousness_text, "life-threatening"),
                ),
            ]
        )

    action = product.get("action_taken") or ""
    action_paragraph = _find_paragraph(document, "☐ Drug withdrawn")
    if action_paragraph:
        action_paragraph.text = "     ".join(
            [
                _checked("Drug withdrawn", action == "Drug withdrawn"),
                _checked("Dose increased", action == "Dose increased"),
                _checked("Dose reduced", action == "Dose reduced"),
                _checked("Dose not changed", action == "Dose not changed"),

            ]
        )

    outcome = case.get("event_outcome") or ""
    outcome_one = _find_paragraph(document, "☐ Recovered")
    if outcome_one:
        outcome_one.text = "     ".join(
            [
                _checked("Recovered", outcome == "Recovered/resolved"),
                _checked("Recovering", outcome == "Recovering/resolving"),
                _checked("Continuing", outcome == "Not recovered/not resolved"),
                _checked(
                    "Recovered with sequelae",
                    outcome == "Recovered with sequelae",
                ),
            ]
        )

    outcome_two = _find_paragraph(document, "☐ Not recovered")
    if outcome_two:
        outcome_two.text = "     ".join(
            [
                _checked(
                    "Not recovered",
                    outcome == "Not recovered/not resolved",
                ),
                _checked("Death", outcome == "Fatal"),
                _checked("Unknown", outcome == "Unknown"),
            ]
        )

    causality = case.get("causality_assessment") or ""
    causality_paragraph = _find_paragraph(document, "☐ Certain")
    if causality_paragraph:
        causality_paragraph.text = "     ".join(
            [
                _checked("Certain", causality == "Certain"),
                _checked(
                    "Probable / Likely",
                    causality == "Probable / Likely",
                ),
                _checked("Possible", causality == "Possible"),
                _checked("Unlikely", causality == "Unlikely"),
                _checked(
                    "Unclassifiable",
                    causality == "Unassessable / Unclassifiable",
                ),
            ]
        )

    reporter_paragraph = _find_paragraph(document, "Name of Reporter:")
    if reporter_paragraph:
        reporter_paragraph.text = (
            f"Name of Reporter: {_value(case.get('reporter_name'))}     "
            f"Designation: {_value(case.get('reporter_profession'))}"
        )

    reporting_date_paragraph = _find_paragraph(document, "Date of Reporting:")
    if reporting_date_paragraph:
        contact = case.get("reporter_email") or case.get("reporter_phone")
        reporting_date_paragraph.text = (
            f"Date of Reporting: {_date(case.get('received_date'))}     "
            f"Email Address / Contact: {_value(contact)}"
        )

    admin_paragraph = _find_paragraph(document, "Report title:")
    if admin_paragraph:
        admin_paragraph.text = (
            f"Report title: {_value(case.get('report_title'))}     "
            f"Form ID number: {case["case_number"]}"
        )

    output_folder = Path(current_app.config["UPLOAD_ROOT"]).resolve() / "reports"
    output_folder.mkdir(parents=True, exist_ok=True)

    output_path = output_folder / f"ADR_{case['case_number']}.docx"
    document.save(output_path)

    return output_path