"""Fill a Word (.docx) template with data and tidy the result.

Template markers (typed anywhere in the Word file, including tables,
headers and footers):

    {{case.case_number}}            a value
    {{#events}} ... {{/events}}     a block repeated for each item of a list,
                                    shown once for a filled value, or removed
                                    when the value is empty. The markers must
                                    each be alone in their own paragraph.
                                    Inside, use {{events.term}} or {{term}}.
    a table row containing {{products.product_name}}
                                    is repeated once per product.

Tidying, so nothing unnecessary is printed:
* a line whose only data is empty (e.g. "Patient phone: {{case.patient_phone}}")
  is removed;
* a table row whose data is all empty is removed;
* a repeated table with no items is removed;
* a heading whose section ends up empty is removed;
* runs of blank paragraphs are collapsed to one.

Formatting of the template (fonts, styles, tables, logos, headers and
footers) is kept; only the markers are replaced.
"""

import copy
import re
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn

MARKER = re.compile(r"\{\{\s*([#/]?)\s*([A-Za-z_][\w.]*)\s*\}\}")
BLOCK_LINE = re.compile(r"^\s*\{\{\s*([#/])\s*([A-Za-z_][\w.]*)\s*\}\}\s*$")
EMPTY = "⁠⁣⁠"  # invisible marker for "no data here"
LABEL_MAX = 90  # a line this short with no data is a label, and is removed


class TemplateError(ValueError):
    """The template cannot be used; the message says why."""


# --------------------------------------------------------------------------
# Values
# --------------------------------------------------------------------------

def lookup(context, name):
    value = context
    for part in name.split("."):
        if isinstance(value, dict):
            value = value.get(part)
        else:
            return None
    return value


def format_value(value):
    """Text for a marker, or EMPTY when there is nothing to print."""
    if value is None:
        return EMPTY
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, datetime):
        return value.strftime("%d %b %Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d %b %Y")
    if isinstance(value, Decimal):
        value = format(value.normalize(), "f")
    if isinstance(value, float):
        value = f"{value:g}"
    if isinstance(value, (list, tuple)):
        parts = [format_value(v) for v in value]
        parts = [p for p in parts if p != EMPTY]
        return ", ".join(parts) if parts else EMPTY
    if isinstance(value, dict):
        return EMPTY
    text = str(value).replace("\r\n", "\n").strip()
    return text if text else EMPTY


# --------------------------------------------------------------------------
# Walking the document
# --------------------------------------------------------------------------

W_P, W_TBL, W_TR, W_TC = qn("w:p"), qn("w:tbl"), qn("w:tr"), qn("w:tc")


def _text(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def _runs(p_el):
    return [r for r in p_el.iter(qn("w:r"))]


def _run_text(r_el):
    return "".join(t.text or "" for t in r_el.iter(qn("w:t")))


def _set_run_text(r_el, text):
    """Replace a run's text, keeping its formatting; newlines become breaks."""
    for child in list(r_el):
        if child.tag in (qn("w:t"), qn("w:br"), qn("w:tab"), qn("w:cr")):
            r_el.remove(child)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        if i:
            r_el.append(r_el.makeelement(qn("w:br"), {}))
        if line:
            t = r_el.makeelement(qn("w:t"), {})
            t.text = line
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            r_el.append(t)


def _block_children(parent):
    return [c for c in parent if c.tag in (W_P, W_TBL)]


def _has_drawing(el):
    return any(True for _ in el.iter(qn("w:drawing"))) or any(True for _ in el.iter(qn("w:pict")))


class _Tracker:
    def __init__(self):
        self.empty_paragraphs = set()   # paragraphs whose every marker was empty
        self.filled_paragraphs = set()  # paragraphs with at least one value
        self.repeat_tables = set()      # tables that had a repeated row


# --------------------------------------------------------------------------
# Replacing markers in a paragraph
# --------------------------------------------------------------------------

def _fill_paragraph(p_el, context, tracker):
    runs = _runs(p_el)
    texts = [_run_text(r) for r in runs]
    full = "".join(texts)
    matches = [m for m in MARKER.finditer(full) if not m.group(1)]
    if not matches:
        return
    starts, pos = [], 0
    for t in texts:
        starts.append(pos)
        pos += len(t)

    def run_at(offset):
        for i in range(len(runs) - 1, -1, -1):
            if starts[i] <= offset and (offset < starts[i] + len(texts[i]) or i == len(runs) - 1):
                return i
        return 0

    any_filled = False
    for m in reversed(matches):
        value = format_value(lookup(context, m.group(2)))
        if value != EMPTY:
            any_filled = True
        a, b = m.start(), m.end()
        first, last = run_at(a), run_at(b - 1)
        if first == last:
            t = texts[first]
            texts[first] = t[: a - starts[first]] + value + t[b - starts[first]:]
        else:
            texts[first] = texts[first][: a - starts[first]] + value
            for i in range(first + 1, last):
                texts[i] = ""
            texts[last] = texts[last][b - starts[last]:]
    for r, t in zip(runs, texts):
        _set_run_text(r, t)
    (tracker.filled_paragraphs if any_filled else tracker.empty_paragraphs).add(p_el)


# --------------------------------------------------------------------------
# Blocks, repeated rows, containers
# --------------------------------------------------------------------------

def _item_context(context, name, item):
    """Context inside a block or repeated row: the item is reachable as
    {{name.field}} and, for convenience, as {{field}}."""
    child = dict(context)
    if isinstance(item, dict):
        child.update(item)
    child[name.split(".")[-1]] = item
    child[name] = item
    return child


def _has_value(value):
    if isinstance(value, dict):
        return any(_has_value(v) for v in value.values())
    if isinstance(value, (list, tuple)):
        return bool(value)
    return format_value(value) != EMPTY


def _fill_container(parent, context, tracker):
    children = _block_children(parent)
    i = 0
    while i < len(children):
        el = children[i]
        if el.tag == W_P:
            line = BLOCK_LINE.match(_text(el))
            if line and line.group(1) == "#":
                name = line.group(2)
                depth, j = 1, i + 1
                while j < len(children):
                    other = BLOCK_LINE.match(_text(children[j])) if children[j].tag == W_P else None
                    if other and other.group(2) == name:
                        depth += 1 if other.group(1) == "#" else -1
                        if depth == 0:
                            break
                    j += 1
                if j >= len(children):
                    raise TemplateError(f"“{{{{#{name}}}}}” has no matching “{{{{/{name}}}}}”.")
                segment = children[i + 1:j]
                value = lookup(context, name)
                if isinstance(value, (list, tuple)):
                    for item in value:
                        copies = [copy.deepcopy(s) for s in segment]
                        holder = parent.makeelement(qn("w:body"), {})
                        for c in copies:
                            holder.append(c)
                        _fill_container(holder, _item_context(context, name, item), tracker)
                        for c in list(holder):
                            el.addprevious(c)
                    for s in segment:
                        parent.remove(s)
                elif _has_value(value) and value is not False:
                    holder = parent.makeelement(qn("w:body"), {})
                    for s in segment:
                        parent.remove(s)
                        holder.append(s)
                    _fill_container(holder, _item_context(context, name, value) if isinstance(value, dict) else context, tracker)
                    for c in list(holder):
                        el.addprevious(c)
                else:
                    for s in segment:
                        parent.remove(s)
                parent.remove(el)
                parent.remove(children[j])
                children = _block_children(parent)
                # restart from where the block was
                i = max(0, i - 1)
                continue
            _fill_paragraph(el, context, tracker)
        elif el.tag == W_TBL:
            _fill_table(el, context, tracker)
        i += 1


def _row_lists(tr, context):
    names = {m.group(2) for m in MARKER.finditer(_text(tr)) if not m.group(1)}
    for name in sorted(names):
        if "." in name:
            prefix = name.rsplit(".", 1)[0]
            if isinstance(lookup(context, prefix), (list, tuple)):
                return prefix
    return None


def _fill_table(tbl, context, tracker):
    for tr in [c for c in tbl if c.tag == W_TR]:
        list_name = _row_lists(tr, context)
        if list_name:
            tracker.repeat_tables.add(tbl)
            for item in lookup(context, list_name) or []:
                new = copy.deepcopy(tr)
                tr.addprevious(new)
                item_ctx = _item_context(context, list_name, item)
                for tc in new.iter(W_TC):
                    _fill_container(tc, item_ctx, tracker)
            tbl.remove(tr)
            continue
        for tc in [c for c in tr if c.tag == W_TC]:
            _fill_container(tc, context, tracker)


# --------------------------------------------------------------------------
# Tidying
# --------------------------------------------------------------------------

def _heading_level(p_el, styles):
    style = p_el.find(qn("w:pPr") + "/" + qn("w:pStyle"))
    if style is not None:
        name = (styles.get(style.get(qn("w:val"))) or style.get(qn("w:val")) or "").lower()
        m = re.match(r"heading\s*(\d)", name)
        if m:
            return int(m.group(1))
        if name == "title":
            return 0
    return None


def _style_names(document):
    names = {}
    try:
        for style in document.styles:
            names[style.style_id] = style.name
    except Exception:
        pass
    return names


def _paragraph_is_blank(p_el):
    if _text(p_el).replace(EMPTY, "").strip() or _has_drawing(p_el):
        return False
    for br in p_el.iter(qn("w:br")):
        if br.get(qn("w:type")) == "page":
            return False
    return True


def _tidy_container(parent, tracker, marked_headings, styles, top=False):
    # 1. Lines whose only data was empty.
    for p in [c for c in parent if c.tag == W_P]:
        if p in tracker.empty_paragraphs and p not in tracker.filled_paragraphs:
            label = _text(p).replace(EMPTY, "").strip()
            if len(label) <= LABEL_MAX and not _has_drawing(p):
                parent.remove(p)
                continue
    # 2. Tables: rows with only empty data; repeated tables with no rows.
    for tbl in [c for c in parent if c.tag == W_TBL]:
        rows = [r for r in tbl if r.tag == W_TR]
        for tr in rows:
            paras = list(tr.iter(W_P))
            had_data = [p for p in paras if p in tracker.empty_paragraphs or p in tracker.filled_paragraphs]
            if had_data and not any(p in tracker.filled_paragraphs for p in had_data):
                tbl.remove(tr)
        rows = [r for r in tbl if r.tag == W_TR]
        if tbl in tracker.repeat_tables and len(rows) <= 1:
            only_header = not rows or not any(p in tracker.filled_paragraphs for p in rows[0].iter(W_P))
            if only_header:
                parent.remove(tbl)
                continue
        if not rows:
            parent.remove(tbl)
            continue
        for tc in tbl.iter(W_TC):
            _tidy_container(tc, tracker, marked_headings, styles)
            if not any(c.tag in (W_P, W_TBL) for c in tc):
                tc.append(tc.makeelement(W_P, {}))
    # 3. Headings whose section is now empty (top level only).
    if top:
        changed = True
        while changed:
            changed = False
            blocks = _block_children(parent)
            for idx, el in enumerate(blocks):
                if el.tag != W_P or el not in marked_headings:
                    continue
                level = _heading_level(el, styles)
                level = 9 if level is None else level
                section = []
                for nxt in blocks[idx + 1:]:
                    nxt_level = _heading_level(nxt, styles) if nxt.tag == W_P else None
                    if nxt_level is not None and nxt_level <= level:
                        break
                    if level == 9 and nxt.tag == W_P and (nxt in marked_headings or _is_bold_heading(nxt)):
                        break
                    section.append(nxt)
                if all(s.tag == W_P and _paragraph_is_blank(s) for s in section):
                    for s in section:
                        parent.remove(s)
                    parent.remove(el)
                    changed = True
                    break
                has_data = any(
                    p in tracker.filled_paragraphs or _has_drawing(p)
                    for s in section for p in ([s] if s.tag == W_P else s.iter(W_P))
                )
                if not has_data:
                    # No data left under this heading; static lines that
                    # follow (e.g. a signature line) stay, the heading and
                    # its empty labels go.
                    for s in section:
                        if s.tag == W_P and (_paragraph_is_blank(s) or _text(s).strip().endswith(":")):
                            parent.remove(s)
                    parent.remove(el)
                    changed = True
                    break
    # 4. Collapse runs of blank paragraphs.
    previous_blank = False
    for el in _block_children(parent):
        blank = el.tag == W_P and _paragraph_is_blank(el) and el.find(qn("w:pPr") + "/" + qn("w:sectPr")) is None
        if blank and previous_blank:
            parent.remove(el)
            continue
        previous_blank = blank


_E = re.escape(EMPTY)
_LEFTOVERS = (
    # "(onset <empty>)" -> ""
    (re.compile(r"\s*[\(\[][^()\[\]]*" + _E + r"[^()\[\]]*[\)\]]"), ""),
    # "A, <empty>, B" / "A · <empty>" -> "A, B" / "A"
    (re.compile(r"\s*[,;·|/–-]\s*" + _E), ""),
    (re.compile(_E + r"\s*[,;·|/–-]\s*"), ""),
)


def _strip_empty_markers(root):
    """Remove the invisible empty markers, with brackets or separators that
    only surrounded missing data."""
    for t in root.iter(qn("w:t")):
        if t.text and EMPTY in t.text:
            text = t.text
            for pattern, replacement in _LEFTOVERS:
                text = pattern.sub(replacement, text)
            t.text = text.replace(EMPTY, "")


_NUMBERED = re.compile(r"^(\s*)(\d+)((?:\.\d+)*)([.)]?\s+)")


def _renumber_headings(body, styles, original_numbers):
    """Close gaps in typed heading numbers ("1., 2., 4." -> "1., 2., 3.")
    left by removed sections. Only top-level numbers are changed, and only
    when the template's own numbering was consecutive."""
    if not original_numbers or original_numbers != list(range(original_numbers[0], original_numbers[0] + len(original_numbers))):
        return
    mapping, next_number = {}, original_numbers[0]
    headings = [el for el in _block_children(body) if el.tag == W_P and _heading_level(el, styles) is not None]
    for el in headings:
        m = _NUMBERED.match(_text(el))
        if m and not m.group(3):
            old = int(m.group(2))
            if old not in mapping:
                mapping[old] = next_number
                next_number += 1
    if all(k == v for k, v in mapping.items()):
        return
    for el in headings:
        runs = _runs(el)
        if not runs:
            continue
        first = _run_text(runs[0])
        m = _NUMBERED.match(first)
        if m and int(m.group(2)) in mapping:
            _set_run_text(runs[0], m.group(1) + str(mapping[int(m.group(2))]) + first[m.end(1) + len(m.group(2)):])


def _heading_numbers(body, styles):
    numbers = []
    for el in _block_children(body):
        if el.tag == W_P and _heading_level(el, styles) is not None:
            m = _NUMBERED.match(_text(el))
            if m and not m.group(3):
                numbers.append(int(m.group(2)))
    return numbers


def _marked_headings(body, styles):
    """Headings (styled, or short bold lines) followed by template data."""
    marked = set()
    blocks = _block_children(body)
    for idx, el in enumerate(blocks):
        if el.tag != W_P:
            continue
        level = _heading_level(el, styles)
        is_bold_label = level is None and _is_bold_heading(el)
        if level is None and not is_bold_label:
            continue
        level = 9 if level is None else level
        for nxt in blocks[idx + 1:]:
            nxt_level = _heading_level(nxt, styles) if nxt.tag == W_P else None
            if nxt_level is not None and nxt_level <= level:
                break
            if level == 9 and nxt.tag == W_P and _is_bold_heading(nxt):
                break
            if MARKER.search(_text(nxt)):
                marked.add(el)
                break
    return marked


def _is_bold_heading(p_el):
    text = _text(p_el).strip()
    if not text or len(text) > 80 or MARKER.search(text) or text.endswith(":"):
        return False
    runs = [r for r in _runs(p_el) if _run_text(r).strip()]
    if not runs:
        return False
    for r in runs:
        rpr = r.find(qn("w:rPr"))
        b = rpr.find(qn("w:b")) if rpr is not None else None
        if b is None or b.get(qn("w:val")) in ("0", "false"):
            return False
    return True


# --------------------------------------------------------------------------
# Public
# --------------------------------------------------------------------------

def _parts(document):
    """Body plus every header and footer, as XML containers."""
    yield document.element.body, True
    seen = set()
    for section in document.sections:
        for part in (section.header, section.footer, section.first_page_header,
                     section.first_page_footer, section.even_page_header, section.even_page_footer):
            try:
                el = part._element
            except Exception:
                continue
            if id(el) in seen:
                continue
            seen.add(id(el))
            yield el, False


def fill_template(template_bytes, context):
    """Return the filled, tidied document as bytes. Raises TemplateError."""
    try:
        document = Document(BytesIO(template_bytes))
    except Exception as error:
        raise TemplateError(f"The file is not a readable Word document ({error}).") from error
    styles = _style_names(document)
    tracker = _Tracker()
    marked = _marked_headings(document.element.body, styles)
    numbers = _heading_numbers(document.element.body, styles)
    for container, top in _parts(document):
        _fill_container(container, context, tracker)
    for container, top in _parts(document):
        _tidy_container(container, tracker, marked, styles, top=top)
        _strip_empty_markers(container)
    _renumber_headings(document.element.body, styles, numbers)
    leftover = [m.group(0) for c, _ in _parts(document) for m in MARKER.finditer(_text(c))]
    if leftover:
        raise TemplateError("Some markers could not be filled: " + ", ".join(sorted(set(leftover))[:10]))
    out = BytesIO()
    document.save(out)
    return out.getvalue()


def template_markers(template_bytes):
    """All markers in the template: (values, block names, problems)."""
    try:
        document = Document(BytesIO(template_bytes))
    except Exception as error:
        raise TemplateError(f"The file is not a readable Word document ({error}).") from error
    values, blocks, problems = set(), set(), []
    for container, _ in _parts(document):
        stack = []
        for el in container.iter(W_P):
            text = _text(el)
            for m in MARKER.finditer(text):
                kind, name = m.group(1), m.group(2)
                if not kind:
                    values.add(name)
                    continue
                if not BLOCK_LINE.match(text):
                    problems.append(
                        f"“{m.group(0)}” must be on a line of its own (its own paragraph)."
                    )
                if kind == "#":
                    stack.append(name)
                    blocks.add(name)
                elif not stack or stack[-1] != name:
                    problems.append(f"“{m.group(0)}” closes a block that was not opened just before it.")
                else:
                    stack.pop()
            if "{{" in text and not MARKER.search(text):
                problems.append(f"Marker not understood near “{text.strip()[:60]}”.")
        for name in stack:
            problems.append(f"“{{{{#{name}}}}}” is never closed with “{{{{/{name}}}}}”.")
    return values, blocks, problems


def document_outline(template_bytes, limit=400):
    """The template's text, numbered, for the AI to read (no markers needed)."""
    document = Document(BytesIO(template_bytes))
    lines = []
    body = document.element.body
    p_index = t_index = 0
    for el in _block_children(body):
        if el.tag == W_P:
            text = _text(el).strip()
            if text:
                lines.append({"ref": f"p{p_index}", "text": text[:200]})
            p_index += 1
        else:
            for r_i, tr in enumerate([r for r in el if r.tag == W_TR]):
                for c_i, tc in enumerate([c for c in tr if c.tag == W_TC]):
                    text = _text(tc).strip()
                    lines.append({"ref": f"t{t_index}r{r_i}c{c_i}", "text": text[:200]})
            t_index += 1
        if len(lines) >= limit:
            break
    return lines


def insert_markers(template_bytes, placements):
    """Write markers into a template that has none.

    placements: [{"ref": "p12" | "t0r3c1", "field": "case.case_number",
                  "how": "after" | "replace" | "row"}]
    "after" adds the marker after the label text, "replace" puts it in place
    of the cell's text, "row" marks a table row to repeat for a list (the
    field is "list.sub" and goes into that cell).
    Returns new template bytes.
    """
    document = Document(BytesIO(template_bytes))
    body = document.element.body
    paragraphs, tables, p_index, t_index = {}, {}, 0, 0
    for el in _block_children(body):
        if el.tag == W_P:
            paragraphs[f"p{p_index}"] = el
            p_index += 1
        else:
            tables[t_index] = el
            t_index += 1

    def cell(ref):
        m = re.fullmatch(r"t(\d+)r(\d+)c(\d+)", ref)
        if not m:
            return None
        tbl = tables.get(int(m.group(1)))
        if tbl is None:
            return None
        rows = [r for r in tbl if r.tag == W_TR]
        r = int(m.group(2))
        if r >= len(rows):
            return None
        cells = [c for c in rows[r] if c.tag == W_TC]
        c = int(m.group(3))
        return cells[c] if c < len(cells) else None

    def write(p_el, marker, how):
        runs = _runs(p_el)
        if how == "replace" or not runs:
            for r in runs[1:]:
                _set_run_text(r, "")
            if runs:
                _set_run_text(runs[0], marker)
            else:
                r = p_el.makeelement(qn("w:r"), {})
                p_el.append(r)
                _set_run_text(r, marker)
            return
        last = runs[-1]
        text = _run_text(last)
        _set_run_text(last, text + ("" if text.endswith(" ") else " ") + marker)

    for placement in placements:
        ref, field, how = placement["ref"], placement["field"], placement.get("how", "after")
        marker = "{{" + field + "}}"
        if ref.startswith("p"):
            p_el = paragraphs.get(ref)
            if p_el is not None:
                write(p_el, marker, "replace" if how == "replace" else "after")
            continue
        tc = cell(ref)
        if tc is None:
            continue
        ps = [p for p in tc if p.tag == W_P]
        if not ps:
            ps = [tc.makeelement(W_P, {})]
            tc.append(ps[0])
        write(ps[0], marker, "replace" if how in ("replace", "row") or not _text(tc).strip() else "after")
    out = BytesIO()
    document.save(out)
    return out.getvalue()
