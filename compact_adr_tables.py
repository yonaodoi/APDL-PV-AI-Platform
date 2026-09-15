from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


template_path = Path(
    r"controlled_templates\8.1ADR REPORTING FORM.docx"
)

document = Document(template_path)


def set_cell_width(cell, width):
    cell.width = width

    cell_properties = cell._tc.get_or_add_tcPr()
    cell_width = cell_properties.find(qn("w:tcW"))

    if cell_width is None:
        cell_width = OxmlElement("w:tcW")
        cell_properties.append(cell_width)

    cell_width.set(qn("w:type"), "dxa")
    cell_width.set(qn("w:w"), str(width))


def format_table(table, total_width, proportions, font_size):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    proportion_total = sum(proportions)

    for column_index, proportion in enumerate(proportions):
        width = int(total_width * proportion / proportion_total)

        for row in table.rows:
            if column_index < len(row.cells):
                set_cell_width(row.cells[column_index], width)

    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)

                if row_index == 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

                for run in paragraph.runs:
                    run.font.size = Pt(font_size)

                    if row_index == 0:
                        run.bold = True


product_table = document.tables[0]

product_headers = [
    "Generic",
    "Brand",
    "Batch",
    "Route / Dose / Frequency / Strength",
    "Start",
    "Stop",
    "Indication",
    "Suspect",
]

for cell, header in zip(product_table.rows[0].cells, product_headers):
    cell.text = header

section = document.sections[0]
product_width = (
    section.page_width
    - section.left_margin
    - section.right_margin
)

format_table(
    product_table,
    product_width,
    proportions=[1.25, 1.10, 0.80, 1.75, 0.70, 0.70, 1.00, 0.40],
    font_size=7,
)


signature_table = None

for table in document.tables[1:]:
    if table.rows:
        first_row = " ".join(
            cell.text.strip()
            for cell in table.rows[0].cells
        )

        if "Signatory" in first_row:
            signature_table = table
            break

if signature_table is None:
    raise RuntimeError("Signatory table was not found.")

signature_table._element.getparent().remove(
    signature_table._element
)

new_table = document.add_table(rows=1, cols=4)
new_table.style = "Table Grid"

headers = [
    "Signatory",
    "Name and designation",
    "Signature",
    "Date",
]

for cell, header in zip(new_table.rows[0].cells, headers):
    cell.text = header

signatories = [
    (
        "Prepared by",
        "AMEKO CHARLES\nDEPUTY Q.P.P.V.",
    ),
    (
        "Reviewed by",
        "YONA ODOI\nQ.P.P.V.",
    ),
    (
        "Authorised by",
        "KEITH ARUHO\nGROUP HEAD, REGULATORY AFFAIRS AND QUALITY",
    ),
]

for role, person_details in signatories:
    row = new_table.add_row()

    for cell, value in zip(
        row.cells,
        [role, person_details, "", ""],
    ):
        cell.text = value

format_table(
    new_table,
    Inches(5.2),
    proportions=[1.2, 2.3, 1.05, 0.65],
    font_size=8,
)

document.save(template_path)

print("ADR tables compacted successfully.")