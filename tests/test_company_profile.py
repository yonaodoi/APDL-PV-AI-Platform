from datetime import date
from io import BytesIO

from docx import Document

from app.services import company_profile

XYZ = company_profile.normalise({
    **company_profile.APDL_PROFILE,
    "legal_name": "Xyz Pharma Limited",
    "letter_name": "Xyz Pharma Ltd",
    "short_name": "XYZ",
    "platform_name": "XYZ Safety",
    "department": "Drug Safety Department",
    "pv_phone": "+254700000000",
    "pv_email": "drugsafety@xyz.example",
    "case_prefix": "xyz-icsr-",
    "signal_prefix": "XYZ-SIG",
    "complaint_prefix": "XYZ-PC",
    "psur_prefix": "XYZ-PSUR",
    "sign_prepared_name": "JANE DOE",
    "sign_prepared_title": "PV OFFICER",
    "sign_reviewed_name": "JOHN ROE",
    "sign_reviewed_title": "QPPV",
    "sign_authorised_name": "MARY POE",
    "sign_authorised_title": "HEAD OF QUALITY",
    "form_code": "XYZ/PV/001",
})


def _use(monkeypatch, profile):
    monkeypatch.setattr(company_profile, "company", lambda: profile)


def _docx_text(data):
    document = Document(BytesIO(data.getvalue() if hasattr(data, "getvalue") else data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts += [c.text for c in row.cells]
    for section in document.sections:
        parts += [p.text for p in section.header.paragraphs]
    return "\n".join(parts)


def test_fresh_install_is_neutral_and_apdl_profile_is_unchanged():
    fresh = company_profile.normalise({})
    assert fresh["short_name"] == "Company" and fresh["sign_reviewed_name"] == ""
    assert "APDL" not in str(fresh) and "Abacus" not in str(fresh)
    apdl = company_profile.normalise(company_profile.APDL_PROFILE)
    assert apdl["case_prefix"] == "APDL-ICSR" and apdl["sign_reviewed_name"] == "YONA ODOI"


def test_values_are_tidied_and_problems_explained():
    assert XYZ["case_prefix"] == "XYZ-ICSR"
    cleared = company_profile.normalise({**XYZ, "sign_prepared_name": "", "address": ""})
    assert cleared["sign_prepared_name"] == "" and cleared["short_name"] == "XYZ"
    assert company_profile.normalise({**XYZ, "brand_colour": "green"})["brand_colour"] == "#145a40"
    problems = company_profile.problems({**XYZ, "short_name": "", "pv_email": "nope", "brand_colour": "#12"})
    assert any("Short name" in p for p in problems)
    assert any("email" in p for p in problems)
    assert any("colour" in p for p in problems)


def test_examples_file_names_and_signatories(monkeypatch):
    _use(monkeypatch, XYZ)
    assert company_profile.example_number("case", date(2026, 1, 1)) == "XYZ-ICSR-26-001"
    assert company_profile.file_prefix() == "XYZ"
    assert company_profile.platform() == "XYZ Safety platform"
    assert company_profile.signatory_rows(2)[1] == ("Reviewed by", "JOHN ROE", "QPPV", "", "")


def test_follow_up_email_and_form_use_the_company(monkeypatch):
    _use(monkeypatch, XYZ)
    from app.services.follow_up_request import build_follow_up_request_docx, email_body

    case = {"case_id": 1, "case_number": "XYZ-ICSR-26-001", "reporter_name": "Dr A",
            "received_date": date(2026, 10, 1)}
    items = [{"number": 1, "label": "Outcome of the reaction", "code": "x", "question": "What was the outcome?", "answer": "text"}]
    body = email_body(case, {"product_name": "DRUG"}, items, date(2026, 10, 20))
    assert "Xyz Pharma Ltd" in body and "+254700000000" in body and "drugsafety@xyz.example" in body
    assert "Abacus" not in body
    text = _docx_text(build_follow_up_request_docx(case, {"product_name": "DRUG"}, items, date(2026, 10, 20)))
    assert "XYZ PHARMA LTD" in text and "DRUG SAFETY DEPARTMENT" in text and "XYZ/PV/001" in text
    assert "JOHN ROE" in text and "MARY POE" in text
    assert "ABACUS" not in text and "ODOI" not in text and "APDL" not in text


def test_without_a_phone_the_email_leaves_it_out(monkeypatch):
    _use(monkeypatch, {**XYZ, "pv_phone": ""})
    from app.services.follow_up_request import email_body

    case = {"case_id": 1, "case_number": "X", "reporter_name": None, "received_date": None}
    body = email_body(case, None, [{"number": 1, "label": "Dates", "code": "x"}], None)
    assert "call us on" not in body and "Tel:" not in body


def test_psur_titles_and_own_document_type_follow_the_company(monkeypatch):
    _use(monkeypatch, XYZ)
    from app.psur.section_definitions import PSUR_SECTION_CHOICES, PSUR_SECTION_TITLES
    from app.services.case_documents import document_type_label, rsi_document_type
    from app.services.rsi_assessment import is_apdl_document

    assert PSUR_SECTION_TITLES["appendix_sponsored_studies"].endswith("all XYZ sponsored studies")
    assert ("appendix_sponsored_studies", "Appendix 3: Listing of all XYZ sponsored studies") in list(PSUR_SECTION_CHOICES)
    assert rsi_document_type("rsi_apdl") == "XYZ Product Information"
    assert document_type_label("rsi_apdl") == "RSI – XYZ product information"
    assert is_apdl_document({"document_type": "XYZ Product Information"})
    assert is_apdl_document({"document_type": "APDL Product Information"})   # older records


def _client(monkeypatch, role="System Administrator"):
    from app import create_app
    from config import TestingConfig

    monkeypatch.setattr("app.services.case_follow_up_reminders.get_open_reminder_count", lambda: 0)
    client = create_app(TestingConfig).test_client()
    with client.session_transaction() as s:
        s["user_id"], s["full_name"], s["role"] = 1, "Admin", role
    return client


def test_pages_show_the_company_name(monkeypatch):
    _use(monkeypatch, {**XYZ, "brand_colour": "#7a1f3d"})
    monkeypatch.setattr(company_profile, "load_profile", lambda: XYZ)
    client = _client(monkeypatch)
    page = client.get("/administration/company-profile").get_data(as_text=True)
    assert "XYZ Safety" in page and "Xyz Pharma Limited" in page and "JOHN ROE" in page
    assert "--forest-mid: #7a1f3d" in page                                     # brand colour applied
    assert "APDL" not in page and "Abacus" not in page

    with client.session_transaction() as s:
        s["role"] = "PV Officer"
    assert client.get("/administration/company-profile").status_code in (302, 403)


def test_saving_records_the_changes(monkeypatch):
    saved, audit = {}, []
    monkeypatch.setattr(company_profile, "load_profile", lambda: company_profile.normalise(company_profile.APDL_PROFILE))
    monkeypatch.setattr(company_profile, "save_profile", lambda values, user_id: saved.update(company_profile.normalise(values)) or saved)
    monkeypatch.setattr("app.administration.routes.write_audit_log", lambda **kw: audit.append(kw))
    client = _client(monkeypatch)
    form = {k: v for k, v in company_profile.APDL_PROFILE.items() if isinstance(v, str)}
    form.update({"short_name": "XYZ", "csrf_token": "x"})
    response = client.post("/administration/company-profile", data=form)
    assert response.status_code == 302
    assert saved["short_name"] == "XYZ"
    assert "Short name: APDL -> XYZ" in audit[0]["details"]

    bad = client.post("/administration/company-profile", data={**form, "legal_name": ""})
    assert bad.status_code == 200 and "Nothing was saved" in bad.get_data(as_text=True)
