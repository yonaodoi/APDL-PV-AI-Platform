from datetime import date

from docx import Document

import app.cases.review_routes as review_routes
from app import create_app
from app.services.follow_up_request import (
    build_follow_up_request_docx,
    build_request_items,
    email_body,
    email_subject,
    group_tasks_by_case,
)
from config import TestingConfig


CASE = {
    "case_id": 9,
    "case_number": "APDL-ICSR-26-009",
    "received_date": date(2026, 9, 12),
    "reporter_name": "Dr Okello",
    "reporter_email": "okello@example.com",
    "patient_initials": "JK",
    "patient_age_years": 34,
    "patient_sex": "Female",
    "event_description": "Generalised urticaria after infusion",
    "follow_up_due_date": None,
}
PRODUCT = {"product_name": "ABPARA"}
CHECKS = [
    {"code": "event_onset_date", "label": "Event onset date", "status": "Review",
     "message": "Confirm the event onset date."},
    {"code": "event_outcome", "label": "Event outcome", "status": "Review",
     "message": "Record the outcome or select Unknown."},
    {"code": "follow_up_due_date", "label": "Follow-up plan", "status": "Review",
     "message": "Set a due date when follow-up is required."},
    {"code": "event_coding", "label": "Event coding", "status": "Review",
     "message": "Code the event."},
    {"code": "fatal_outcome_consistency", "label": "Fatal outcome consistency", "status": "Review",
     "message": "Death is recorded as a seriousness criterion but the outcome is Recovered/resolved. Confirm which is correct."},
    {"code": "reported_event", "label": "Reported event", "status": "Pass", "message": "Recorded."},
]


def test_items_skip_internal_and_passed_checks_and_are_numbered():
    items = build_request_items(CASE, PRODUCT, CHECKS)

    assert [i["code"] for i in items] == [
        "event_onset_date", "event_outcome", "fatal_outcome_consistency",
    ]
    assert [i["number"] for i in items] == [1, 2, 3]
    assert items[0]["answer"] == "date"
    assert "Generalised urticaria after infusion" in items[0]["question"]
    assert items[1]["answer"] == "options"
    assert "Unknown" in items[1]["options"]
    assert "do not agree" in items[2]["question"]
    assert items[2]["question"].count("which is correct") == 1


def test_word_form_has_context_questions_and_tick_boxes():
    items = build_request_items(CASE, PRODUCT, CHECKS)
    output = build_follow_up_request_docx(
        CASE, PRODUCT, items, date(2026, 10, 15), prepared_by="Charles Ameko",
        today=date(2026, 10, 8),
    )
    document = Document(output)
    text = "\n".join(p.text for p in document.paragraphs)
    cells = "\n".join(c.text for t in document.tables for r in t.rows for c in r.cells)

    assert "CASE FOLLOW-UP FORM" in text
    assert "1. When the reaction started" in text
    assert "2. Outcome of the reaction" in text
    assert "☐  Recovered/resolved" in text
    assert "Follow-up plan" not in text and "Event coding" not in text
    assert "APDL-ICSR-26-009" in cells
    assert "ABPARA" in cells
    assert "15 Oct 2026" in cells
    assert "CHARLES AMEKO" in cells


def test_email_text_lists_the_questions_once():
    items = build_request_items(CASE, PRODUCT, CHECKS)
    body = email_body(CASE, PRODUCT, items, date(2026, 10, 15))

    assert email_subject(CASE) == "Follow-up request for adverse reaction report APDL-ICSR-26-009"
    assert body.startswith("Dear Dr Okello,")
    assert "  1. When the reaction started" in body
    assert "  3. Outcome and seriousness" in body
    assert "by 15 Oct 2026" in body


def test_grouping_separates_reporter_and_internal_tasks():
    tasks = [
        {"case_id": 1, "case_number": "A", "task_id": 1, "check_code": "event_outcome",
         "due_date": date(2026, 10, 20), "reporter_email": "x@example.com"},
        {"case_id": 1, "case_number": "A", "task_id": 2, "check_code": "event_coding",
         "due_date": date(2026, 10, 18), "reporter_email": "x@example.com"},
        {"case_id": 2, "case_number": "B", "task_id": 3, "check_code": "reported_event",
         "due_date": date(2026, 10, 10), "reporter_email": ""},
    ]
    groups = group_tasks_by_case(tasks, {1: date(2026, 10, 1)})

    assert [g["case_number"] for g in groups] == ["B", "A"]
    a = groups[1]
    assert len(a["reporter_tasks"]) == 1 and len(a["internal_tasks"]) == 1
    assert a["due_date"] == date(2026, 10, 18)
    assert a["can_email"] and a["last_request"] == date(2026, 10, 1)
    assert not groups[0]["can_email"]


class _Cursor:
    def __init__(self, log):
        self.log = log
        self.rowcount = 0

    def execute(self, sql, parameters=()):
        self.log.append((" ".join(sql.split()), parameters))

    def fetchone(self):
        return {"case_id": 9}


class _Transaction:
    def __init__(self, log):
        self.log = log

    def __enter__(self):
        return _Cursor(self.log)

    def __exit__(self, *args):
        return False


OPEN_TASKS = [
    {"task_id": 21, "check_code": "event_onset_date", "due_date": date(2026, 10, 15)},
    {"task_id": 22, "check_code": "event_outcome", "due_date": date(2026, 10, 15)},
    {"task_id": 23, "check_code": "event_coding", "due_date": date(2026, 10, 15)},
]


def _client(monkeypatch, case):
    import app.services.follow_up_automation as automation

    review_checks = [c for c in CHECKS if c["status"] == "Review"]
    monkeypatch.setattr(
        automation, "load_case_request",
        lambda case_id: (case, PRODUCT, review_checks, OPEN_TASKS),
    )
    monkeypatch.setattr(review_routes, "load_case_request", automation.load_case_request)
    monkeypatch.setattr(
        "app.services.case_follow_up_reminders.get_open_reminder_count", lambda: 0
    )
    app = create_app(TestingConfig)
    client = app.test_client()
    with client.session_transaction() as user_session:
        user_session["user_id"] = 1
        user_session["full_name"] = "Charles Ameko"
        user_session["role"] = "System Administrator"
    return client


def test_download_returns_one_word_form_for_the_case(monkeypatch):
    response = _client(monkeypatch, CASE).get("/cases/9/follow-up-request.docx")

    assert response.status_code == 200
    assert "follow-up-request-APDL-ICSR-26-009.docx" in response.headers["Content-Disposition"]


def test_send_emails_once_and_records_each_reporter_task(monkeypatch):
    import app.services.follow_up_automation as automation

    client = _client(monkeypatch, CASE)
    sent, log = [], []
    monkeypatch.setattr(
        "app.services.follow_up_email.send_follow_up_email",
        lambda app, recipient, subject, body, attachment, filename: sent.append((recipient, subject, filename)),
    )
    monkeypatch.setattr(automation, "transaction", lambda: _Transaction(log))
    monkeypatch.setattr("app.services.case_follow_up.sync_follow_up_flags", lambda cursor, case_id: 2)
    monkeypatch.setattr("app.services.case_workflow.auto_move_status", lambda *args: True)

    response = client.post("/cases/9/follow-up-request/send")

    assert response.status_code == 302
    assert "#case-9" in response.headers["Location"]
    assert sent == [("okello@example.com",
                     "Follow-up request for adverse reaction report APDL-ICSR-26-009",
                     "follow-up-request-APDL-ICSR-26-009.docx")]
    deliveries = [p for sql, p in log if "INSERT INTO pv.case_follow_up_email_deliveries" in sql]
    assert [(p[0], p[3], p[4], p[5]) for p in deliveries] == [
        (21, 1, "Request", False), (22, 1, "Request", False),
    ]
    assert any("Follow-up request emailed" in str(p) for sql, p in log)


def test_send_refuses_without_reporter_email(monkeypatch):
    client = _client(monkeypatch, {**CASE, "reporter_email": None})
    sent = []
    monkeypatch.setattr("app.services.follow_up_email.send_follow_up_email", lambda *a: sent.append(a))

    response = client.post("/cases/9/follow-up-request/send")

    assert response.status_code == 302
    assert sent == []
