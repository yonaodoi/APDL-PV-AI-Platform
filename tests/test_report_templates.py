from datetime import date
from io import BytesIO

from docx import Document

from app.services import docx_fill
from app.services.report_fields import REPORT_TYPES, sample_context, unknown_markers


def _doc(build):
    d = Document()
    build(d)
    out = BytesIO()
    d.save(out)
    return out.getvalue()


def _texts(data):
    d = Document(BytesIO(data))
    lines = []
    for el in d.element.body:
        tag = el.tag.split("}")[1]
        if tag == "p":
            lines.append("".join(t.text or "" for t in el.iter(docx_fill.qn("w:t"))))
        elif tag == "tbl":
            from docx.table import Table

            lines.append([[c.text for c in r.cells] for r in Table(el, d).rows])
    return lines


def _icsr(d):
    d.sections[0].header.paragraphs[0].text = "Case {{case.case_number}}"
    p = d.add_paragraph()
    p.add_run("Case number: ").bold = True
    p.add_run("{{case.")
    p.add_run("case_number}}")
    d.add_heading("1. Patient", 1)
    d.add_paragraph("Initials: {{case.patient_initials}}")
    d.add_paragraph("Phone: {{case.patient_phone}}")
    d.add_heading("2. Laboratory", 1)
    d.add_paragraph("{{case.laboratory_results}}")
    d.add_heading("3. Products", 1)
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Product", "Batch"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "{{products.product_name}}", "{{products.batch_number}}"
    d.add_heading("4. Events", 1)
    d.add_paragraph("{{#events}}")
    d.add_paragraph("- {{term}} (coded {{events.preferred_term}})")
    d.add_paragraph("{{/events}}")
    d.add_paragraph("Prepared by: ________")


CONTEXT = {
    "case": {"case_number": "APDL-ICSR-26-001", "patient_initials": "A.N.", "patient_phone": None,
             "laboratory_results": "", "received_date": date(2026, 10, 1)},
    "products": [{"product_name": "ABPARA", "batch_number": "B1"}, {"product_name": "AMOXIL", "batch_number": None}],
    "events": [{"term": "Rash", "preferred_term": "Rash"}, {"term": "Swelling", "preferred_term": None}],
}


def test_fills_values_repeats_rows_and_tidies_empty_parts():
    out = docx_fill.fill_template(_doc(_icsr), CONTEXT)
    lines = _texts(out)

    assert "Case number: APDL-ICSR-26-001" in lines
    assert "Initials: A.N." in lines
    assert not any(isinstance(l, str) and l.startswith("Phone") for l in lines)      # empty line removed
    assert not any(isinstance(l, str) and "Laboratory" in l for l in lines)          # empty section removed
    assert "2. Products" in lines and "3. Events" in lines                           # numbering closed up
    assert [["Product", "Batch"], ["ABPARA", "B1"], ["AMOXIL", ""]] in lines
    assert "- Rash (coded Rash)" in lines
    assert "- Swelling" in lines                                                    # empty brackets removed
    assert "Prepared by: ________" in lines
    header = Document(BytesIO(out)).sections[0].header.paragraphs[0].text
    assert header == "Case APDL-ICSR-26-001"


def test_empty_list_removes_its_table_and_heading():
    out = docx_fill.fill_template(_doc(_icsr), {**CONTEXT, "products": [], "events": []})
    lines = _texts(out)
    assert not any(isinstance(l, list) for l in lines)
    assert not any(isinstance(l, str) and ("Products" in l or "Events" in l) for l in lines)


def test_markers_are_listed_and_problems_found():
    values, blocks, problems = docx_fill.template_markers(_doc(_icsr))
    assert "case.patient_initials" in values and "events" in blocks and problems == []

    def broken(d):
        d.add_paragraph("{{#events}}")
        d.add_paragraph("text {{/products}}")
    _, _, problems = docx_fill.template_markers(_doc(broken))
    assert any("own paragraph" in p for p in problems)
    assert any("never closed" in p or "not opened" in p for p in problems)


def test_unknown_markers_are_reported():
    assert unknown_markers("case_report", {"case.case_number", "case.patient_nmae", "products.batch_number", "term"}) == ["case.patient_nmae"]


def test_every_report_type_sample_fills_a_template_using_all_its_markers():
    for report_type, spec in REPORT_TYPES.items():
        def build(d, spec=spec):
            for name in spec["fields"]:
                d.add_paragraph(f"{name}: {{{{{name}}}}}")
            for list_name, info in spec["lists"].items():
                t = d.add_table(rows=2, cols=len(info["fields"]))
                for i, sub in enumerate(info["fields"]):
                    t.rows[0].cells[i].text = sub
                    t.rows[1].cells[i].text = f"{{{{{list_name}.{sub}}}}}"
        out = docx_fill.fill_template(_doc(build), sample_context(report_type))
        assert b"{{" not in out or "{{" not in "".join(str(l) for l in _texts(out)), report_type


def test_ai_placements_are_written_as_markers():
    def form(d):
        d.add_paragraph("Patient initials:")
        t = d.add_table(rows=2, cols=2)
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Product", "Batch"
    template = _doc(form)
    outline = docx_fill.document_outline(template)
    assert {"ref": "p0", "text": "Patient initials:"} in outline

    from app.services.template_mapping import clean_placements

    raw = [
        {"ref": "p0", "field": "case.patient_initials", "how": "after"},
        {"ref": "t0r1c0", "field": "products.product_name", "how": "row"},
        {"ref": "t0r1c1", "field": "products.batch_number", "how": "row"},
        {"ref": "p99", "field": "case.case_number"},                 # unknown place
        {"ref": "t0r0c0", "field": "case.not_a_field"},              # unknown field
    ]
    placements = clean_placements("case_report", raw, outline + [{"ref": "t0r1c0", "text": ""}, {"ref": "t0r1c1", "text": ""}])
    assert [p["ref"] for p in placements] == ["p0", "t0r1c0", "t0r1c1"]
    marked = docx_fill.insert_markers(template, placements)
    out = docx_fill.fill_template(marked, CONTEXT)
    lines = _texts(out)
    assert "Patient initials: A.N." in lines
    assert [["Product", "Batch"], ["ABPARA", "B1"], ["AMOXIL", ""]] in lines


def test_check_template_requires_markers_and_known_fields(monkeypatch):
    from app import create_app
    from app.services.report_templates import check_template
    from config import TestingConfig

    app = create_app(TestingConfig)
    with app.app_context():
        good = check_template("case_report", _doc(_icsr))
        assert good["passed"] is True and good["problems"] == []
        plain = check_template("case_report", _doc(lambda d: d.add_paragraph("Patient initials:")))
        assert plain["passed"] is False and "no markers" in plain["problems"][0]
        typo = check_template("case_report", _doc(lambda d: d.add_paragraph("{{case.patient_nmae}}")))
        assert typo["passed"] is False and "patient_nmae" in typo["problems"][0]


def test_active_template_is_used_and_failure_falls_back(monkeypatch):
    from app import create_app
    from app.services import report_templates
    from config import TestingConfig

    app = create_app(TestingConfig)
    template = _doc(_icsr)
    monkeypatch.setattr(report_templates, "active_template",
                        lambda report_type: {"template_id": 1, "name": "ICSR v2", "stored_filename": "x.docx"})
    monkeypatch.setattr(report_templates, "read_template", lambda row, original=False: template)
    with app.app_context():
        output, notice = report_templates.render_with_active("case_report", lambda: CONTEXT)
        assert notice is None and b"PK" == output.getvalue()[:2]

        def broken():
            raise RuntimeError("database is down")
        output, notice = report_templates.render_with_active("case_report", broken)
        assert output is None and "standard layout was used" in notice


def test_template_pages_load(monkeypatch):
    from app import create_app
    from config import TestingConfig

    monkeypatch.setattr("app.services.case_follow_up_reminders.get_open_reminder_count", lambda: 0)
    monkeypatch.setattr("app.services.report_templates.list_templates", lambda: [])
    client = create_app(TestingConfig).test_client()
    with client.session_transaction() as s:
        s["user_id"], s["full_name"], s["role"] = 1, "Yona Odoi", "System Administrator"
    page = client.get("/report-templates/").get_data(as_text=True)
    assert "Individual case report (ICSR / ADR form)" in page and "Upload and test" in page
    guide = client.get("/report-templates/guide/case_report").get_data(as_text=True)
    assert "{{products.batch_number}}" in guide

    with client.session_transaction() as s:
        s["role"] = "PV Officer"
    assert client.get("/report-templates/").status_code in (302, 403)
