import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


template_path = Path(
    r"controlled_templates\8.1ADR REPORTING FORM.docx"
)

backup_path = Path(
    r"controlled_templates\8.1ADR REPORTING FORM.original.docx"
)

if not template_path.exists():
    raise FileNotFoundError(
        f"Template not found: {template_path.resolve()}"
    )

if not backup_path.exists():
    shutil.copy2(template_path, backup_path)

document = Document(template_path)

for section in document.sections:
    footer = section.footer

    for table in list(footer.tables):
        table._element.getparent().remove(table._element)

    for paragraph in footer.paragraphs:
        paragraph.text = ""

document.add_paragraph("")

heading = document.add_paragraph()
heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

heading_run = heading.add_run("Signatories")
heading_run.bold = True
heading_run.font.size = Pt(12)

table = document.add_table(rows=1, cols=5)
table.style = "Table Grid"

headers = [
    "Signatory",
    "Name",
    "Designation",
    "Signature",
    "Date",
]

for cell, value in zip(table.rows[0].cells, headers):
    cell.text = value
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

signatories = [
    (
        "Prepared by",
        "AMEKO CHARLES",
        "DEPUTY Q.P.P.V.",
    ),
    (
        "Reviewed by",
        "YONA ODOI",
        "Q.P.P.V.",
    ),
    (
        "Authorised by",
        "KEITH ARUHO",
        "GROUP HEAD OF REGULATORY AFFAIRS AND QUALITY",
    ),
]

for role, name, designation in signatories:
    row = table.add_row()

    values = [
        role,
        name,
        designation,
        "",
        "",
    ]

    for cell, value in zip(row.cells, values):
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

document.save(template_path)

print("ADR template updated successfully.")
print(f"Backup created at: {backup_path}")